"""Research Specification Word Renderer — generates a professionally formatted
Microsoft Word (.docx) document from an approved Research Specification.

Uses python-docx.  Applies Hunter Intelligence branding with deep violet
(#5B2C9D) headings and professional typography.
"""
from __future__ import annotations

import json
import logging
import os
import re
import time
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.shared import Inches, Pt, Cm, RGBColor

from . import intelligence_store as store
from . import config

logger = logging.getLogger(__name__)

BRAND_VIOLET = RGBColor(0x5B, 0x2C, 0x9D)
DARK_TEXT = RGBColor(0x33, 0x33, 0x33)
LIGHT_TEXT = RGBColor(0x66, 0x66, 0x66)


def render_spec_docx(spec_id: int) -> str:
    """Render a Research Specification as a Word document.

    Returns the file path to the generated .docx.
    """
    spec_row = store.get_spec_by_id(spec_id)
    if not spec_row:
        raise ValueError(f"Spec {spec_id} not found")

    spec = spec_row["spec"]
    project_name = spec.get("project_name", "Research Specification")
    sections = spec.get("sections", {})
    section_order = spec.get("section_order", list(store.SPEC_SECTION_ORDER))
    section_titles = spec.get("section_titles", {})

    doc = Document()
    _setup_document(doc)
    _add_cover_page(doc, project_name, spec_row)
    _add_table_of_contents(doc, section_order, section_titles, sections)

    for key in section_order:
        section = sections.get(key)
        if not section:
            continue
        title = section.get("title", section_titles.get(key, key.replace("_", " ").title()))
        content = section.get("content", "")

        doc.add_page_break()
        _add_section_heading(doc, title)
        _render_section_content(doc, key, content)

        if section.get("edited"):
            _add_edited_marker(doc)

    _add_header_footer(doc, project_name)

    output_dir = Path(config.DATA_DIR) / "specifications"
    output_dir.mkdir(parents=True, exist_ok=True)

    safe_name = re.sub(r'[^\w\s-]', '', project_name).strip().replace(' ', '_')[:50]
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"Research_Specification_{safe_name}_{timestamp}.docx"
    filepath = str(output_dir / filename)

    doc.save(filepath)
    store.update_spec_docx(spec_id, filepath)
    logger.info("Rendered specification %d to %s", spec_id, filepath)
    return filepath


# ─── Document setup ───────────────────────────────────────────────────────────

def _setup_document(doc: Document):
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)
    font.color.rgb = DARK_TEXT

    for level in range(1, 4):
        name = f"Heading {level}"
        if name in doc.styles:
            hs = doc.styles[name]
            hs.font.color.rgb = BRAND_VIOLET
            hs.font.bold = True
            hs.font.name = "Calibri"
            hs.font.size = Pt(18 - (level * 2))


def _add_cover_page(doc: Document, project_name: str, spec_row: dict):
    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("HUNTER INTELLIGENCE")
    run.font.size = Pt(12)
    run.font.color.rgb = BRAND_VIOLET
    run.font.bold = True
    run.font.name = "Calibri"

    _add_violet_divider(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Research Specification")
    run.font.size = Pt(28)
    run.font.color.rgb = BRAND_VIOLET
    run.font.bold = True
    run.font.name = "Calibri"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(project_name)
    run.font.size = Pt(16)
    run.font.color.rgb = DARK_TEXT
    run.font.name = "Calibri"

    _add_violet_divider(doc)

    meta_lines = [
        f"Version: {spec_row.get('version', 1)}",
        f"Status: {spec_row.get('approval_status', 'pending').replace('_', ' ').title()}",
        f"Generated: {datetime.now().strftime('%B %d, %Y')}",
    ]
    if spec_row.get("generation_source"):
        meta_lines.append(f"Source: {spec_row['generation_source'].title()}")

    for line in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        run.font.size = Pt(10)
        run.font.color.rgb = LIGHT_TEXT
        run.font.name = "Calibri"


def _add_table_of_contents(
    doc: Document, section_order: list, section_titles: dict, sections: dict
):
    doc.add_page_break()
    h = doc.add_heading("Table of Contents", level=1)

    for i, key in enumerate(section_order, 1):
        sec = sections.get(key, {})
        title = sec.get("title", section_titles.get(key, key.replace("_", " ").title()))
        p = doc.add_paragraph()
        run = p.add_run(f"{i}. {title}")
        run.font.size = Pt(11)
        run.font.color.rgb = DARK_TEXT
        run.font.name = "Calibri"


def _add_section_heading(doc: Document, title: str):
    h = doc.add_heading(title, level=1)
    _add_violet_divider(doc)


def _add_violet_divider(doc: Document):
    p = doc.add_paragraph()
    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "5B2C9D")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_edited_marker(doc: Document):
    p = doc.add_paragraph()
    run = p.add_run("[Analyst Edited]")
    run.font.size = Pt(8)
    run.font.color.rgb = BRAND_VIOLET
    run.font.italic = True


def _add_header_footer(doc: Document, title: str):
    for section in doc.sections:
        header = section.header
        hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = hp.add_run(f"Hunter Intelligence — {title}")
        run.font.size = Pt(8)
        run.font.color.rgb = LIGHT_TEXT
        run.font.name = "Calibri"

        footer = section.footer
        fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run("CONFIDENTIAL — For internal research use only")
        run.font.size = Pt(7)
        run.font.color.rgb = LIGHT_TEXT
        run.font.name = "Calibri"


# ─── Section content rendering ────────────────────────────────────────────────

def _render_section_content(doc: Document, key: str, content: Any):
    if isinstance(content, str):
        _render_text_content(doc, content)
    elif isinstance(content, list):
        _render_list_content(doc, key, content)
    elif isinstance(content, dict):
        _render_dict_content(doc, key, content)
    else:
        _render_text_content(doc, str(content))


def _render_text_content(doc: Document, text: str):
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("**") and line.endswith("**"):
            p = doc.add_paragraph()
            run = p.add_run(line.strip("*"))
            run.bold = True
            run.font.name = "Calibri"
            run.font.color.rgb = DARK_TEXT
        elif ":" in line and len(line.split(":")[0]) < 40:
            label, value = line.split(":", 1)
            p = doc.add_paragraph()
            run = p.add_run(f"{label.strip()}: ")
            run.bold = True
            run.font.name = "Calibri"
            run.font.color.rgb = DARK_TEXT
            run = p.add_run(value.strip())
            run.font.name = "Calibri"
            run.font.color.rgb = DARK_TEXT
        else:
            p = doc.add_paragraph(line)


def _render_list_content(doc: Document, key: str, items: list):
    if not items:
        doc.add_paragraph("No items defined.", style="List Bullet")
        return

    first = items[0]
    if isinstance(first, str):
        for item in items:
            doc.add_paragraph(str(item), style="List Bullet")
    elif isinstance(first, dict):
        _render_dict_list(doc, key, items)
    else:
        for item in items:
            doc.add_paragraph(str(item), style="List Bullet")


def _render_dict_list(doc: Document, key: str, items: list[dict]):
    if key == "research_questions":
        _render_research_questions_table(doc, items)
    elif key == "entities":
        _render_entities_table(doc, items)
    elif key in ("exclusions", "risks", "assumptions"):
        _render_generic_table(doc, items)
    elif key == "deliverables":
        _render_deliverables_table(doc, items)
    elif key == "metrics":
        _render_metrics_table(doc, items)
    elif key == "question_method_mapping":
        _render_mapping_table(doc, items)
    elif key == "clarifications":
        _render_clarifications_list(doc, items)
    elif key == "dependencies":
        _render_dependencies_table(doc, items)
    else:
        for item in items:
            parts = []
            for k, v in item.items():
                parts.append(f"{k}: {v}")
            doc.add_paragraph(" | ".join(parts), style="List Bullet")


def _render_dict_content(doc: Document, key: str, data: dict):
    if key == "scope_dimensions":
        _render_scope_dimensions(doc, data)
    elif key == "audiences":
        _render_audiences(doc, data)
    elif key == "inclusions":
        _render_inclusions(doc, data)
    elif key == "methodology":
        _render_methodology_detail(doc, data)
    elif key == "data_requirements":
        _render_data_requirements(doc, data)
    elif key == "approval_status":
        _render_approval_status(doc, data)
    else:
        for k, v in data.items():
            label = k.replace("_", " ").title()
            p = doc.add_paragraph()
            run = p.add_run(f"{label}: ")
            run.bold = True
            run.font.name = "Calibri"
            if isinstance(v, list):
                run2 = p.add_run(", ".join(str(i) for i in v))
            else:
                run2 = p.add_run(str(v))
            run2.font.name = "Calibri"


# ─── Typed renderers ──────────────────────────────────────────────────────────

def _render_research_questions_table(doc: Document, rqs: list[dict]):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    headers = ["ID", "Question", "Priority", "Source"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for rq in rqs:
        row = table.add_row()
        row.cells[0].text = str(rq.get("question_id", ""))
        row.cells[1].text = str(rq.get("question", ""))
        row.cells[2].text = str(rq.get("priority", "")).capitalize()
        row.cells[3].text = str(rq.get("source", "")).capitalize()


def _render_entities_table(doc: Document, entities: list[dict]):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    headers = ["Entity", "Type", "Confidence", "Reasoning"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for ent in entities:
        row = table.add_row()
        row.cells[0].text = str(ent.get("name", ""))
        row.cells[1].text = str(ent.get("type", "")).capitalize()
        row.cells[2].text = str(ent.get("confidence", "")).capitalize()
        row.cells[3].text = str(ent.get("reasoning", ""))


def _render_generic_table(doc: Document, items: list[dict]):
    if not items:
        return
    keys = list(items[0].keys())
    table = doc.add_table(rows=1, cols=len(keys))
    table.style = "Light Grid Accent 1"
    for i, k in enumerate(keys):
        table.rows[0].cells[i].text = k.replace("_", " ").title()
    for item in items:
        row = table.add_row()
        for i, k in enumerate(keys):
            row.cells[i].text = str(item.get(k, ""))


def _render_deliverables_table(doc: Document, items: list[dict]):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    headers = ["Deliverable", "Format", "Status"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for item in items:
        row = table.add_row()
        row.cells[0].text = str(item.get("deliverable", ""))
        row.cells[1].text = str(item.get("format", ""))
        row.cells[2].text = str(item.get("status", "planned")).capitalize()


def _render_metrics_table(doc: Document, items: list[dict]):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    headers = ["Metric", "Target", "Description"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for item in items:
        row = table.add_row()
        row.cells[0].text = str(item.get("metric", ""))
        row.cells[1].text = str(item.get("target", ""))
        row.cells[2].text = str(item.get("description", ""))


def _render_mapping_table(doc: Document, items: list[dict]):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    headers = ["Question ID", "Question", "Method"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for item in items:
        row = table.add_row()
        row.cells[0].text = str(item.get("question_id", ""))
        row.cells[1].text = str(item.get("question", ""))
        row.cells[2].text = str(item.get("method", ""))


def _render_clarifications_list(doc: Document, items: list[dict]):
    for item in items:
        q = item.get("question", "")
        blocking = not item.get("can_proceed_without", True)
        p = doc.add_paragraph()
        run = p.add_run(f"{'[BLOCKING] ' if blocking else ''}{q}")
        run.bold = blocking
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00) if blocking else DARK_TEXT
        if item.get("why_it_matters"):
            p2 = doc.add_paragraph()
            run2 = p2.add_run(f"  Impact: {item['why_it_matters']}")
            run2.font.size = Pt(9)
            run2.font.color.rgb = LIGHT_TEXT
            run2.font.name = "Calibri"


def _render_dependencies_table(doc: Document, items: list[dict]):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    headers = ["Dependency", "Required For", "Status"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for item in items:
        row = table.add_row()
        row.cells[0].text = str(item.get("dependency", ""))
        row.cells[1].text = str(item.get("required_for", ""))
        row.cells[2].text = str(item.get("status", "pending")).capitalize()


def _render_scope_dimensions(doc: Document, data: dict):
    for key in ["platforms", "geography", "time_period", "languages", "content_types"]:
        val = data.get(key, "")
        label = key.replace("_", " ").title()
        p = doc.add_paragraph()
        run = p.add_run(f"{label}: ")
        run.bold = True
        run.font.name = "Calibri"
        if isinstance(val, list):
            run2 = p.add_run(", ".join(str(v) for v in val) if val else "Not specified")
        else:
            run2 = p.add_run(str(val) if val else "Not specified")
        run2.font.name = "Calibri"


def _render_audiences(doc: Document, data: dict):
    primary = data.get("primary_audience", "")
    if primary:
        p = doc.add_paragraph()
        run = p.add_run("Primary Audience: ")
        run.bold = True
        run.font.name = "Calibri"
        p.add_run(str(primary)).font.name = "Calibri"

    desc = data.get("description", "")
    if desc:
        doc.add_paragraph(desc)

    segments = data.get("segments", [])
    if segments:
        doc.add_heading("Audience Segments", level=2)
        for seg in segments:
            if isinstance(seg, dict):
                name = seg.get("name", seg.get("segment_id", ""))
                role = seg.get("role", "")
                doc.add_paragraph(f"{name} — {role}" if role else name, style="List Bullet")


def _render_inclusions(doc: Document, data: dict):
    for key in ["brands", "platforms", "expected_analyses", "deliverables"]:
        items = data.get(key, [])
        if items:
            label = key.replace("_", " ").title()
            doc.add_heading(label, level=2)
            for item in items:
                doc.add_paragraph(str(item), style="List Bullet")


def _render_methodology_detail(doc: Document, data: dict):
    primary = data.get("primary", "")
    if primary:
        p = doc.add_paragraph()
        run = p.add_run("Primary Method: ")
        run.bold = True
        run.font.name = "Calibri"
        p.add_run(primary).font.name = "Calibri"

    reasoning = data.get("reasoning", "")
    if reasoning:
        p = doc.add_paragraph()
        run = p.add_run("Rationale: ")
        run.bold = True
        run.font.name = "Calibri"
        p.add_run(reasoning).font.name = "Calibri"

    alternatives = data.get("alternatives_considered", [])
    if alternatives:
        doc.add_heading("Alternatives Considered", level=2)
        for alt in alternatives:
            if isinstance(alt, dict):
                doc.add_paragraph(
                    f"{alt.get('method', '')} — {alt.get('why_less_appropriate', '')}",
                    style="List Bullet",
                )
            else:
                doc.add_paragraph(str(alt), style="List Bullet")


def _render_data_requirements(doc: Document, data: dict):
    for key in ["data_sources", "required_fields", "sample_size", "time_range",
                "geographic_filter", "language_filter"]:
        val = data.get(key, "")
        label = key.replace("_", " ").title()
        p = doc.add_paragraph()
        run = p.add_run(f"{label}: ")
        run.bold = True
        run.font.name = "Calibri"
        if isinstance(val, list):
            p.add_run(", ".join(str(v) for v in val)).font.name = "Calibri"
        else:
            p.add_run(str(val) if val else "Not specified").font.name = "Calibri"


def _render_approval_status(doc: Document, data: dict):
    approved = data.get("specification_approved", False)
    p = doc.add_paragraph()
    run = p.add_run("Specification Approved: ")
    run.bold = True
    run.font.name = "Calibri"
    status_run = p.add_run("Yes" if approved else "No — Pending Review")
    status_run.font.color.rgb = RGBColor(0x00, 0x80, 0x00) if approved else RGBColor(0xCC, 0x80, 0x00)
    status_run.font.name = "Calibri"

    gate = data.get("gate_status", "")
    if gate:
        p2 = doc.add_paragraph()
        run2 = p2.add_run(f"Gate Status: {gate}")
        run2.font.name = "Calibri"
        run2.font.color.rgb = DARK_TEXT
