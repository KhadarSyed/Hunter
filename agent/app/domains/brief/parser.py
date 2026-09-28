"""Client-brief text extraction.

`extract_brief_text()` is what the upload endpoint uses: NVIDIA Nemotron-Parse first
(layout-aware Markdown for PDF / Word / PowerPoint / Excel), falling back to native
extraction when NVIDIA is unavailable or returns nothing. BRIEF_EXTRACTION_MODE=native
skips NVIDIA entirely. `parse_brief()` is the native path, also used by the
Brief-to-Deck watcher.
"""
from __future__ import annotations

import logging
import os
from pathlib import Path

import openpyxl
from docx import Document
from pptx import Presentation

from ...core import nvidia_parse_client
from ...core.nvidia_parse_client import NvidiaParseError

logger = logging.getLogger(__name__)

# Native extraction supports these; NVIDIA additionally handles legacy .doc/.ppt/.xls.
SUPPORTED_SUFFIXES = {".docx", ".txt", ".pptx", ".pdf", ".xlsx"}
UPLOAD_SUFFIXES = SUPPORTED_SUFFIXES | {".doc", ".ppt", ".xls"}


def parse_brief(path: Path) -> str:
    """Native (offline) text extraction."""
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix == ".pptx":
        return _parse_pptx(path)
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".xlsx":
        return _parse_xlsx(path)
    raise ValueError(f"Unsupported brief format: {suffix}")


def extract_brief_text(path: Path) -> dict:
    """Extract a brief's text, NVIDIA first. Returns {text, method, model, pages}.

    method: "nvidia" (Nemotron-Parse) or "native". Blocking (network + Office
    conversion) — call from a worker thread, not the event loop.
    """
    suffix = path.suffix.lower()
    use_nvidia = (os.getenv("BRIEF_EXTRACTION_MODE", "nvidia").lower() != "native"
                  and suffix in nvidia_parse_client.SUPPORTED_EXTENSIONS
                  and nvidia_parse_client.is_reachable())
    if use_nvidia:
        try:
            text, pages = nvidia_parse_client.parse_document(path.read_bytes(), path.name)
            if text.strip():
                return {"text": text, "method": "nvidia", "model": nvidia_parse_client.model_name(), "pages": pages}
            logger.warning("NVIDIA parse returned no text for %s; using native extraction", path.name)
        except NvidiaParseError as e:
            logger.warning("NVIDIA parse failed for %s (%s); using native extraction", path.name, e)
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"{suffix} needs NVIDIA parsing (or LibreOffice/MS Office), which is unavailable")
    return {"text": parse_brief(path), "method": "native", "model": None, "pages": None}


def _parse_docx(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _parse_pptx(path: Path) -> str:
    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        lines = [s.text_frame.text.strip() for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        if lines:
            parts.append(f"[Slide {i}]\n" + "\n".join(lines))
    return "\n\n".join(parts)


def _parse_pdf(path: Path) -> str:
    import pymupdf

    with pymupdf.open(str(path)) as doc:
        return "\n".join(page.get_text() for page in doc).strip()


def _parse_xlsx(path: Path) -> str:
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        sheets = []
        for ws in wb.worksheets:
            rows = ["  ".join(str(c) for c in row if c is not None).strip() for row in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if r]
            if rows:
                sheets.append((f"[Sheet: {ws.title}]\n" if len(wb.worksheets) > 1 else "") + "\n".join(rows))
        return "\n\n".join(sheets)
    finally:
        wb.close()


def guess_client_name(brief_text: str, filename: str) -> str:
    """Best-effort client name guess; refined later by the LLM reasoning pass."""
    stem = Path(filename).stem
    for sep in ("_", "-"):
        if sep in stem:
            token = stem.split(sep)[0].strip()
            if token and token.lower() not in ("brief", "client", "new"):
                return token
    return stem.strip() or "Unknown Client"
