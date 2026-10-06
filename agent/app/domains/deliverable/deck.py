"""Assemble the Baby Skincare reference deck on top of the Hunter base deck."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.opc.packuri import PackURI
from pptx.util import Inches, Pt

from . import blocks, style
from .citations import CitationRegistry
from .metrics import month_label

CITES_PER_PAGE = 18
COVER_SUBTITLE_PT = 22
_EXPERT_LABELS = {"other_hcp": "Other HCP", "non_hcp_expert": "Non-HCP expert"}


LOGO_MAX_W, LOGO_MAX_H = 0.9, 0.45


def _add_logo(slide, path, x: float, y: float, max_h: float) -> None:
    """Fit a logo inside LOGO_MAX_W x max_h inches, preserving its aspect ratio."""
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), height=Inches(max_h))
    if pic.width > Inches(LOGO_MAX_W):
        ratio = Inches(LOGO_MAX_W) / pic.width
        pic.width, pic.height = Inches(LOGO_MAX_W), int(pic.height * ratio)


def expert_label(key: str) -> str:
    return _EXPERT_LABELS.get(key, key.replace("_", " ").capitalize())


def paginate(entries: list, per_page: int) -> list[list]:
    return [entries[i:i + per_page] for i in range(0, len(entries), per_page)] or [[]]


def keep_cover_and_closing(prs) -> None:
    ids = prs.slides._sldIdLst
    for sld in list(ids)[1:-1]:
        prs.part.drop_rel(sld.rId)
        ids.remove(sld)
    # python-pptx names new slides slide{len+1}.xml; renumber survivors so names never collide
    for i, slide in enumerate(prs.slides, start=1):
        slide.part.partname = PackURI(f"/ppt/slides/slide{i}.xml")


def _move_to_end(prs, index: int) -> None:
    ids = prs.slides._sldIdLst
    el = list(ids)[index]
    ids.remove(el)
    ids.append(el)


def _replace_text(shape, lines: list[str]) -> None:
    tf = shape.text_frame
    first = tf.paragraphs[0].runs[0] if tf.paragraphs[0].runs else None
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    if first is None:
        tf.paragraphs[0].add_run().text = lines[0]
    else:
        first.text = lines[0]
        for r in list(tf.paragraphs[0].runs)[1:]:
            r._r.getparent().remove(r._r)
    for line in lines[1:]:
        r = tf.add_paragraph().add_run()
        r.text = line
        if first is not None:
            r.font.name, r.font.size, r.font.bold = first.font.name, first.font.size, False
            if first.font.color and first.font.color.type is not None:
                r.font.color.rgb = first.font.color.rgb


def _set_cover(slide, title: str, subtitle: str, date_label: str) -> None:
    boxes = [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    for sh in boxes:
        txt = sh.text_frame.text.strip()
        low = txt.lower()
        if "©" in txt or "infovision" in low or "invision" in low:
            continue
        if any(ch.isdigit() for ch in txt) and len(txt) < 20:
            _replace_text(sh, [date_label])
    title_box = max((sh for sh in boxes if not any(ch.isdigit() for ch in sh.text_frame.text)
                     and "©" not in sh.text_frame.text), key=lambda sh: sh.width * sh.height, default=None)
    if title_box is not None:
        title_box.width = Inches(style.SLIDE_W) - 2 * title_box.left
        _replace_text(title_box, [title, subtitle] if subtitle else [title])
        if subtitle:
            title_box.text_frame.paragraphs[1].runs[0].font.size = Pt(COVER_SUBTITLE_PT)


def _source(cfg, theme_metrics=None) -> str:
    if theme_metrics and theme_metrics.get("date_min"):
        return (f"SOURCE: MELTWATER  |  {month_label(theme_metrics['date_min'][:7])} – "
                f"{month_label(theme_metrics['date_max'][:7])}")
    return f"SOURCE: MELTWATER  |  {cfg['period_label']}"


def _content(prs, cfg, m, kicker, title, summary, theme_metrics=None):
    s = blocks.new_content_slide(prs)
    blocks.add_header(s, kicker, title, summary)
    blocks.add_footer(s, _source(cfg, theme_metrics), m["base_n"])
    return s


def _summary(ins: list[dict]) -> str:
    return ins[0]["text"] + blocks.citation_suffix(ins[0]["citations"]) if ins else ""


def _peak_callout(t: dict, registry: CitationRegistry, articles_by_url: dict) -> str:
    parts = []
    for p in t["peaks"]:
        cite = next((f" [{registry.cite(articles_by_url[u])}]" for u in p["top_urls"] if u in articles_by_url), "")
        parts.append(f"#{p['rank']} {month_label(p['month'])}: {p['count']}{cite}")
    return "Top-5 peaks — " + "; ".join(parts) if parts else ""


def _trend_slide(prs, cfg, m, key, label, ins, side_bars: list[dict], side_title: str,
                 registry: CitationRegistry, articles_by_url: dict):
    t = m["themes"][key]
    s = _content(prs, cfg, m, f"{cfg['title']} - Editorial", f"{label}: Volume & Peaks", _summary(ins), t)
    cats = [month_label(x["month"]) for x in t["monthly"]]
    vals = [x["count"] for x in t["monthly"]]
    peak_months = {p["month"] for p in t["peaks"]}
    if cats:
        blocks.add_line_chart_with_peaks(s, 0.4, 1.45, 8.2, 2.5, cats, vals,
                                         {i for i, x in enumerate(t["monthly"]) if x["month"] in peak_months})
    blocks.add_text(s, 0.45, 3.98, 8.1, 0.35, _peak_callout(t, registry, articles_by_url), 9, True, style.VIOLET)
    blocks.add_text(s, 7.3, 1.5, 1.3, 0.3, f"N = {t['count']}", 11, True)
    bars = side_bars or [{"name": o["outlet"], "count": o["count"]} for o in t["outlets"][:6]]
    blocks.add_text(s, 8.9, 1.45, 4.0, 0.3, side_title, 10, True, style.VIOLET)
    if bars:
        blocks.add_bar_chart(s, 8.8, 1.75, 4.2, 2.5, [b["name"] for b in bars], [b["count"] for b in bars], style.PEACH)
    blocks.add_insight_cards(s, 0.4, 4.4, 12.5, 2.75, ins[1:4], cols=3)
    return s


def build_deck(cfg, m, cm, ins, registry: CitationRegistry, articles_by_url, logos, gauge_png, out_path) -> Path:
    prs = Presentation(cfg["reference_deck"])
    keep_cover_and_closing(prs)
    _set_cover(prs.slides[0], cfg["title"], cfg.get("subtitle", ""), cfg["date_label"])
    labels = {t["key"]: t["label"] for t in cfg["themes"]}
    kicker = f"{cfg['title']} - Editorial"
    share_line = (f"Share of collected category coverage (N={m['base_n']} unique articles, "
                  f"Meltwater + manual extraction, {cfg['period_label']}).")

    # 2 Contents
    s = _content(prs, cfg, m, kicker, "Contents", "")
    blocks.add_text(s, 0.8, 1.7, 11.5, 5.3, "\n".join([
        "1. Objective & Scope", "2. Executive Summary", "3. Coverage Mix", "4. Deal / Sale-led Coverage",
        "5. Parenting Advice", "6. Expert-led Coverage", "7. Celebrity-led Coverage", "8. Brands in the Conversation",
        "9. Key Takeaways & Implications", "10. Appendix: Methodology & Citations"]), 18, False, "1F1F1F")
    # 3 Objective & scope
    s = _content(prs, cfg, m, kicker, "Objective & Scope", "Earned editorial coverage of the baby skincare category.")
    blocks.add_text(s, 0.6, 1.7, 7.4, 4.0, "Questions from the brief:\n" +
                    "\n".join(f"• {q}" for q in cfg["brief_questions"]), 13, False, "1F1F1F")
    blocks.add_text(s, 8.3, 1.7, 4.6, 4.0, f"Geography: {cfg['geography']}\nTime period: {cfg['period_label']}\n"
                    f"Sources: Meltwater exports + manual extraction\nBase: N={m['base_n']} unique articles "
                    f"({m['multi_theme']} appear in two themes)", 12, True, style.VIOLET)
    # 4 Executive summary
    s = _content(prs, cfg, m, kicker, "Executive Summary", share_line)
    order = ("deal", "parenting", "expert", "celebrity")
    blocks.add_kpi_tiles(s, 0.5, 1.55, 12.3, 2.25, [
        {"value": f"{m['themes'][k]['share']}%", "label": labels[k],
         "note": f"{m['themes'][k]['count']} of {m['base_n']} articles"} for k in order])
    blocks.add_insight_cards(s, 0.5, 3.95, 12.3, 3.2, ins["exec"][:4], cols=2)
    # 5 Coverage mix
    s = _content(prs, cfg, m, kicker, "Coverage Mix", share_line)
    keys = ("celebrity", "deal", "expert", "parenting")
    blocks.add_doughnut(s, 0.4, 1.45, 5.6, 3.0, [labels[k] for k in keys], [m["themes"][k]["share"] for k in keys])
    blocks.add_text(s, 0.4, 4.45, 5.6, 0.3, "Shares can sum to more than 100%: an article can carry two themes.", 8)
    blocks.add_bar_chart(s, 6.4, 1.45, 6.5, 3.0, [labels[k] for k in keys], [m["themes"][k]["count"] for k in keys])
    blocks.add_insight_cards(s, 0.4, 4.85, 12.5, 2.3, ins["mix"][:3], cols=3)
    # 6 Deal-led
    _trend_slide(prs, cfg, m, "deal", labels["deal"], ins["deal"],
                 [{"name": r["retailer"], "count": r["count"]} for r in cm["deal"]["retailers"][:6]], "Top retailers",
                 registry, articles_by_url)
    # 7 Parenting advice
    s = _content(prs, cfg, m, kicker, "Parenting Advice", _summary(ins["parenting"]), m["themes"]["parenting"])
    blocks.add_text(s, 0.5, 1.42, 12, 0.3, f"Small sample: n={m['themes']['parenting']['count']} articles — "
                    "read as directional.", 9, True, style.PEACH)
    rows = cm["parenting"]["rows"][:14]
    if rows:
        blocks.add_table(s, 0.5, 1.8, 7.6, 0.3 * (len(rows) + 1), ["Article", "Topic", "Outlet", "Cite"], rows,
                         [4.2, 1.4, 1.4, 0.6])
    blocks.add_insight_cards(s, 8.4, 1.8, 4.5, 5.3, ins["parenting"][1:3], cols=1)
    # 8 Expert: who is cited
    e = cm["expert"]
    s = _content(prs, cfg, m, kicker, "Expert-led: Who Is Cited", _summary(ins["expert"]), m["themes"]["expert"])
    if e["type_counts"]:
        blocks.add_doughnut(s, 0.4, 1.45, 5.2, 3.2, [expert_label(k) for k in e["type_counts"]],
                            list(e["type_counts"].values()), "0")
    if gauge_png and Path(gauge_png).exists():
        s.shapes.add_picture(str(gauge_png), Inches(5.8), Inches(1.45), Inches(3.4))
    blocks.add_text(s, 5.8, 3.75, 3.4, 0.8, f"Brand-affiliated: {e['affiliated_n']} of {e['affiliation_known_n']} "
                    f"experts with a stated affiliation ({e['affiliation_unknown']} not stated)", 9)
    outlets = m["themes"]["expert"]["outlets"][:6]
    if outlets:
        blocks.add_bar_chart(s, 9.4, 1.45, 3.6, 3.2, [o["outlet"] for o in outlets], [o["count"] for o in outlets],
                             style.MINT)
    blocks.add_insight_cards(s, 0.4, 4.85, 12.5, 2.3, ins["expert"][1:4], cols=3)
    # 9 Expert: named experts
    s = _content(prs, cfg, m, kicker, "Expert-led: Named Experts",
                 "Experts quoted or cited, with their stated brand affiliation.", m["themes"]["expert"])
    named = e.get("named", [])[:16]
    if named:
        blocks.add_table(s, 0.5, 1.55, 12.3, 0.32 * (len(named) + 1), ["Expert", "Type", "Affiliation", "Outlet", "Cite"],
                         named, [3.0, 2.2, 3.0, 3.3, 0.8])
    # 10 Celebrity: volume & peaks
    _trend_slide(prs, cfg, m, "celebrity", labels["celebrity"], ins["celebrity"],
                 [{"name": c["name"], "count": c["count"]} for c in cm["celebrity"]["top"][:6]],
                 "Most-featured celebrities", registry, articles_by_url)
    # 11 Celebrity: sentiment drivers
    t = m["themes"]["celebrity"]
    s = _content(prs, cfg, m, kicker, "Celebrity-led: Sentiment Drivers", _summary(ins["celebrity"]), t)
    if t["sentiment"]:
        sk = [k for k in ("positive", "neutral", "negative") if k in t["sentiment"]]
        total = sum(t["sentiment"][k] for k in sk)
        blocks.add_doughnut(s, 0.4, 1.45, 4.8, 3.2, [k.title() for k in sk],
                            [round(100 * t["sentiment"][k] / total, 1) for k in sk])
    blocks.add_insight_cards(s, 5.4, 1.45, 7.5, 5.7, ins["celebrity"][4:8] or ins["celebrity"][1:3], cols=2)
    # 12 Brands
    s = _content(prs, cfg, m, kicker, "Brands in the Conversation", _summary(ins["brands"]))
    brands = cm["brands"][:8]
    brand_cards = ins["brands"][1:3]
    chart_w = 6.4 if brand_cards else 11.2
    if brands:
        blocks.add_bar_chart(s, 1.6, 1.45, chart_w, 5.6, [b["brand"] for b in brands], [b["count"] for b in brands])
        step = 5.6 / len(brands)
        for i, b in enumerate(brands):
            logo = logos.get(b["brand"])
            if logo and Path(logo).exists():
                _add_logo(s, logo, 0.35, 1.55 + i * step, min(LOGO_MAX_H, step * 0.75))
    blocks.add_insight_cards(s, 8.3, 1.45, 4.6, 5.7, brand_cards, cols=1)
    # 13 Key takeaways
    s = _content(prs, cfg, m, kicker, "Key Takeaways", "")
    blocks.add_insight_cards(s, 0.4, 1.45, 12.5, 5.7, ins["takeaways"][:6], cols=3)
    # 14 Implications
    s = _content(prs, cfg, m, kicker, "Implications & Whitespace", "")
    blocks.add_insight_cards(s, 0.4, 1.45, 12.5, 5.7, ins["implications"][:4], cols=2)
    # 15 Methodology
    s = _content(prs, cfg, m, "Appendix", "Definitions & Methodology", "")
    blocks.add_text(s, 0.6, 1.55, 12, 5.5, "\n\n".join([
        f"Base: union of the four theme exports, de-duplicated by URL — N={m['base_n']} unique articles; "
        f"{m['multi_theme']} articles appear in two themes, so theme shares can sum to more than 100%.",
        "Shares reflect the collected theme exports; each export's breadth depends on its query, so shares are "
        "not a census of all category coverage.",
        "Date ranges differ by theme; each trend slide states its own range. Peaks are the five months with the "
        "most articles.",
        "Expert type, brand affiliation, celebrities, retailers and topics were labelled from article text by an AI "
        "model with a fixed label set; anything the text does not state is reported as 'not stated'.",
        "Every insight cites the numbered articles that support it; see the citation list.",
    ]), 11)
    # 16+ Citations (paginated)
    pages = paginate(registry.entries(), CITES_PER_PAGE)
    for pi, page in enumerate(pages, start=1):
        title = "Citations" + (f" ({pi}/{len(pages)})" if len(pages) > 1 else "")
        s = _content(prs, cfg, m, "Appendix", title, "")
        rows = [[f"[{c['n']}]", c["outlet"][:28], c["date"], c["title"][:70], c["url"][:80]] for c in page]
        if rows:
            blocks.add_table(s, 0.3, 1.5, 12.7, 0.3 * (len(rows) + 1), ["#", "Outlet", "Date", "Headline", "URL"],
                             rows, [0.5, 1.9, 1.0, 4.5, 4.8])
    # the closing slide sits at index 1 after pruning; move it to the end
    _move_to_end(prs, 1)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return out_path
