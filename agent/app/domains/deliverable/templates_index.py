"""Template stage: index the reference decks once and pick the one closest to the project's scope."""
from __future__ import annotations

import logging
import math
import re
from pathlib import Path

from pptx import Presentation

from ...core import config, store

logger = logging.getLogger(__name__)
DEFAULT_TEMPLATE = config.TEMPLATES_DIR / "Hunter PR Research_Johnson’s (Baby) Editorial _May 2026.pptx"
_WORD = re.compile(r"[a-z]{3,}")
MAX_TEXT = 4000


def is_usable(path: Path) -> bool:
    try:
        prs = Presentation(str(path))
    except Exception:
        return False
    return len(prs.slides) >= 2 and any(l.name == "Blank" for l in prs.slide_layouts)


def _deck_text(path: Path) -> str:
    prs = Presentation(str(path))
    parts = [path.stem] + [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
    return " ".join(parts)[:MAX_TEXT]


def _embed(llm, text: str):
    try:
        vec = llm.embed(text) if llm is not None and hasattr(llm, "embed") else None
        return [float(x) for x in vec] if vec is not None else None   # embed clients may return numpy arrays
    except Exception as e:   # embeddings are optional: token overlap still works
        logger.info("template embedding unavailable (%s)", e)
        return None


def index_templates(llm=None) -> int:
    folder = Path(config.TEMPLATES_DIR)
    count = 0
    for path in sorted(folder.glob("*.pptx")) if folder.exists() else []:
        if path.name.startswith("~$") or not is_usable(path):
            continue
        text = _deck_text(path)
        store.save_reference_deck(str(path), text, _embed(llm, text))
        count += 1
    return count


def _tokens(text: str) -> set[str]:
    return set(_WORD.findall(text.lower()))


def _cosine(a, b) -> float:
    dot = sum(x * y for x, y in zip(a, b))
    na, nb = math.sqrt(sum(x * x for x in a)), math.sqrt(sum(y * y for y in b))
    return dot / (na * nb) if na and nb else 0.0


def choose_template(scope_text: str, llm=None) -> Path:
    path = _choose(scope_text, llm)
    if not Path(path).exists():
        raise FileNotFoundError(f"No usable PowerPoint template: add a .pptx to {config.TEMPLATES_DIR}")
    return path


def _choose(scope_text: str, llm=None) -> Path:
    folder = str(Path(config.TEMPLATES_DIR))

    def current():
        return [d for d in store.list_reference_decks() if d["path"].startswith(folder) and Path(d["path"]).exists()]

    decks = current() or (current() if index_templates(llm) else [])
    if not decks:
        return DEFAULT_TEMPLATE
    query_vec = _embed(llm, scope_text) if any(d["embedding"] for d in decks) else None
    query_tokens = _tokens(scope_text)

    def score(d):
        if query_vec and d["embedding"]:
            return _cosine(query_vec, d["embedding"])
        return len(query_tokens & _tokens(d["text"])) / (len(query_tokens) or 1)

    best = max(decks, key=score)
    return Path(best["path"]) if score(best) > 0 else DEFAULT_TEMPLATE
