"""Render stage: one visual deck for any project - cover with brand banner or category photo, executive answers,
overview, per-RQ chart slides with cited insight cards, tables, logos and icons, takeaways, methodology and a
citation appendix - in the brand heading font, accent and chart tints."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from . import blocks, style
from .amcharts_png import render_treemap_png
from .deck import _add_logo, _move_to_end, _set_cover, keep_cover_and_closing, paginate
from .engine_types import RQ, Section
from .gauge import render_gauge_png

CITES_PER_PAGE = 14
CHART_TOP = 1.5
FOOT_SOURCE = "SOURCE: MELTWATER"
EMU_PER_IN = 914400
COVER_IMAGE_TOP = 1.2
COVER_IMAGE_W = 4.2
COVER_IMAGE_MAX_H = 5.0


@dataclass
class DeckInput:
    title: str
    subtitle: str
    date_label: str
    period_label: str
    rqs: list[RQ]
    overview: Section
    sections_by_rq: dict[str, list[Section]]
    insights_by_rq: dict[str, list[dict]]
    answers: list[dict]
    takeaways: list[dict]
    methodology: list[str]
    citations: list[dict]
    base_n: int
    palette: list[str]
    accent: str
    title_font: str
    logos: dict[str, Path]
    icons: dict[str, Path]
    hero: Path | None
    hero_credit: str
    brand_image: Path | None
    flag: Path | None
    rq_titles: dict[str, str] = field(default_factory=dict)   # short slide titles from the analysis plan


def _slide(prs, inp: DeckInput, kicker: str, title: str, summary: str = "", icon: Path | None = None):
    s = blocks.new_content_slide(prs)
    blocks.add_header(s, kicker, title, summary, font=inp.title_font, accent=inp.accent)
    blocks.add_footer(s, f"{FOOT_SOURCE}  |  {inp.period_label}", inp.base_n)
    if icon and icon.exists():
        s.shapes.add_picture(str(icon), Inches(style.SLIDE_W - 0.95), Inches(0.3), height=Inches(0.5))
    return s


def _picture(slide, path: Path, x, y, w, h):
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def _draw_chart(slide, sec: Section, inp: DeckInput, work: Path, x, y, w, h):
    c = sec.chart
    kind, cats, vals = c["kind"], c.get("categories", []), c.get("values", [])
    if kind == "line_peaks":
        blocks.add_line_chart_with_peaks(slide, x, y, w, h, cats, vals, c.get("peaks", []))
    elif kind == "bar":
        fmt = blocks.PERCENT_FORMAT if c.get("unit") == "percent" else "0"
        blocks.add_bar_chart(slide, x, y, w, h, cats, vals, color=inp.palette[0], number_format=fmt)
    elif kind == "column":
        blocks.add_column_chart(slide, x, y, w, h, cats, vals, c.get("peaks", []))
    elif kind == "doughnut":
        blocks.add_doughnut(slide, x, y, w, h, cats, vals, colors=inp.palette)
    elif kind == "gauge":
        png = render_gauge_png(vals[0], cats[0] if cats else "", work / f"{sec.id}-gauge.png")
        _picture(slide, png, x, y, w, min(h, w * 0.62))
    elif kind == "treemap":
        png = render_treemap_png(cats, vals, work / f"{sec.id}-treemap.png", palette=inp.palette)
        _picture(slide, png, x, y, w, min(h, w * 0.62))
    elif kind == "kpi":
        note = sec.facts[0].split(": ", 1)[-1] if sec.facts else ""
        blocks.add_kpi_tiles(slide, x, y, w, min(h, 2.2),
                             [{"value": f"{vals[0]}%", "label": cats[0] if cats else "", "note": note}])


def _cover(prs, inp: DeckInput):
    cover = prs.slides[0]
    _set_cover(cover, inp.title, inp.subtitle, inp.date_label)
    use_brand = bool(inp.brand_image and inp.brand_image.exists())
    image = inp.brand_image if use_brand else inp.hero
    if image and image.exists():
        pic = cover.shapes.add_picture(str(image), Inches(style.SLIDE_W - 4.6), Inches(COVER_IMAGE_TOP),
                                       width=Inches(COVER_IMAGE_W))
        if pic.height > Inches(COVER_IMAGE_MAX_H):          # tall images: fit the height, keep the aspect ratio
            ratio = Inches(COVER_IMAGE_MAX_H) / pic.height
            pic.height, pic.width = Inches(COVER_IMAGE_MAX_H), int(pic.width * ratio)
        if not use_brand and inp.hero_credit:
            blocks.add_text(cover, style.SLIDE_W - 4.6, COVER_IMAGE_TOP + pic.height / EMU_PER_IN + 0.05, 4.2, 0.3,
                            inp.hero_credit, 7)
    if inp.flag and inp.flag.exists():
        cover.shapes.add_picture(str(inp.flag), Inches(0.6), Inches(6.2), height=Inches(0.45))


def _exec_summary(prs, inp: DeckInput):
    s = _slide(prs, inp, "EXECUTIVE SUMMARY", "What the coverage says",
               f"Base: {inp.base_n} unique articles across {len(inp.rqs)} questions", inp.icons.get("overview"))
    tiles = [{"value": a["value"], "label": a["rq_id"], "note": a["answer"]} for a in inp.answers[:4]]
    blocks.add_kpi_tiles(s, style.MARGIN, CHART_TOP, style.SLIDE_W - 2 * style.MARGIN, 2.4, tiles)
    rows = [[a["rq_id"], a["question"][:80], a["answer"]] for a in inp.answers]
    blocks.add_table(s, style.MARGIN, 4.1, style.SLIDE_W - 2 * style.MARGIN, 0.3 * (len(rows) + 1),
                     ["Question", "Asked", "Answer"], rows, [1.0, 6.0, 5.4])


def _overview(prs, inp: DeckInput, work: Path):
    s = _slide(prs, inp, "OVERVIEW", inp.overview.title, inp.overview.facts[0] if inp.overview.facts else "",
               inp.icons.get("overview"))
    _draw_chart(s, inp.overview, inp, work, style.MARGIN, CHART_TOP, 7.4, 4.6)
    blocks.add_text(s, 8.2, CHART_TOP, 4.6, 4.6, "\n".join(inp.overview.facts[1:8]), 10)


def _logo_row(slide, inp: DeckInput, names: list[str], x: float, y: float):
    for j, name in enumerate([n for n in names if n in inp.logos][:6]):
        _add_logo(slide, inp.logos[name], x + (j % 3) * 1.45, y + (j // 3) * 0.75, 0.55)


def _rq_slides(prs, inp: DeckInput, rq: RQ, work: Path):
    sections = inp.sections_by_rq.get(rq.id, [])
    drawable = [x for x in sections if not x.skipped and x.chart and x.chart["kind"] != "kpi"]
    kpi = next((x for x in sections if x.module == "share_kpi" and not x.skipped), None)
    insights = inp.insights_by_rq.get(rq.id, [])
    summary = kpi.facts[0].split(": ", 1)[-1] if kpi else "No articles for this question"
    title = inp.rq_titles.get(rq.id) or rq.question[:60]
    s = _slide(prs, inp, rq.id, title, summary, inp.icons.get("share_kpi"))
    blocks.add_text(s, style.MARGIN, CHART_TOP - 0.12, style.SLIDE_W - 2 * style.MARGIN, 0.2, rq.question, 8)
    if not kpi:
        blocks.add_text(s, style.MARGIN, CHART_TOP + 0.4, 12, 0.6, "No articles for this question", 18, True,
                        inp.accent, font=inp.title_font)
        return
    first = drawable[:2]
    width = (style.SLIDE_W - 2 * style.MARGIN - 0.3) / max(1, len(first))
    for i, sec in enumerate(first):
        x = style.MARGIN + i * (width + 0.3)
        blocks.add_text(s, x, CHART_TOP, width, 0.3, sec.title, 11, True, inp.accent, font=inp.title_font)
        _draw_chart(s, sec, inp, work, x, CHART_TOP + 0.3, width, 2.7)
        notes = [n for n in sec.notes if n != "logos"]
        if notes:
            blocks.add_text(s, x, CHART_TOP + 3.0, width, 0.25, "; ".join(notes), 8)
    if insights:
        blocks.add_insight_cards(s, style.MARGIN, CHART_TOP + 3.35, style.SLIDE_W - 2 * style.MARGIN, 2.0,
                                 insights[:3], cols=3)
    for sec in drawable[2:]:
        s2 = _slide(prs, inp, rq.id, sec.title, sec.facts[0] if sec.facts else "", inp.icons.get(sec.module))
        _draw_chart(s2, sec, inp, work, style.MARGIN, CHART_TOP, 7.6, 4.6)
        blocks.add_text(s2, 8.4, CHART_TOP, 4.4, 3.8, "\n".join(sec.facts[1:7]), 10)
        if "logos" in sec.notes:
            _logo_row(s2, inp, sec.chart["categories"], 8.4, 5.6)
    for sec in (x for x in sections if not x.skipped and x.table):
        s3 = _slide(prs, inp, rq.id, sec.title, sec.facts[0] if sec.facts else "", inp.icons.get(sec.module))
        rows = sec.table["rows"][:12]
        blocks.add_table(s3, style.MARGIN, CHART_TOP, style.SLIDE_W - 2 * style.MARGIN, 0.32 * (len(rows) + 1),
                         sec.table["header"], rows, sec.table["col_widths"])
    skipped = [x for x in sections if x.skipped and x.module != "share_kpi"]
    if skipped:
        blocks.add_text(s, style.MARGIN, 6.55, 12, 0.25,
                        "Not shown: " + "; ".join(f"{x.title} ({x.skipped})" for x in skipped), 8)


def _takeaways(prs, inp: DeckInput):
    s = _slide(prs, inp, "KEY TAKEAWAYS", "What to do next", "", inp.icons.get("takeaway"))
    blocks.add_insight_cards(s, style.MARGIN, CHART_TOP, style.SLIDE_W - 2 * style.MARGIN, 4.8, inp.takeaways[:6],
                             cols=3)


def _methodology(prs, inp: DeckInput):
    s = _slide(prs, inp, "APPENDIX", "Definitions & Methodology", "", inp.icons.get("methodology"))
    blocks.add_text(s, style.MARGIN, CHART_TOP, style.SLIDE_W - 2 * style.MARGIN, 4.8, "\n".join(inp.methodology), 10)


def _citations(prs, inp: DeckInput) -> list[int]:
    pages = paginate(inp.citations, CITES_PER_PAGE)
    numbers = []
    for i, page in enumerate(pages, start=1):
        s = _slide(prs, inp, "APPENDIX", f"Citations ({i}/{len(pages)})")
        rows = [[str(c["n"]), c["outlet"], c["title"][:70], c["date"], c["url"][:60]] for c in page]
        blocks.add_table(s, style.MARGIN, CHART_TOP, style.SLIDE_W - 2 * style.MARGIN, 0.3 * (len(rows) + 1),
                         ["#", "Outlet", "Headline", "Date", "URL"], rows, [0.5, 2.0, 5.2, 1.2, 3.5])
        numbers.append(len(prs.slides))
    return numbers


def build_generic_deck(inp: DeckInput, template: Path, work_dir: Path, out_path: Path) -> tuple[Path, list[int]]:
    work_dir.mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(template))
    keep_cover_and_closing(prs)
    _cover(prs, inp)
    _exec_summary(prs, inp)
    _overview(prs, inp, work_dir)
    for rq in inp.rqs:
        _rq_slides(prs, inp, rq, work_dir)
    _takeaways(prs, inp)
    _methodology(prs, inp)
    appendix = _citations(prs, inp)
    _move_to_end(prs, 1)                  # the template closing slide goes last...
    appendix = [n - 1 for n in appendix]  # ...so every appendix slide moves up by one
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    # Not generated prose, so the fact check skips them: citation appendix, template cover and closing slide
    return out_path, sorted(set(appendix) | {1, len(prs.slides)})
