"""QC stage: layout checks with automatic fixes, and a fact check that every slide figure is a computed fact."""
from __future__ import annotations

import logging
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

from ...core.llm_synthesis import _figures
from . import qc

logger = logging.getLogger(__name__)
MIN_FONT_PT = 7
SHRINK_PT = 2
MAX_FIX_ROUNDS = 3
_CITATION = re.compile(r"\[\d+\]")
_YEAR = re.compile(r"\b(19|20)\d{2}\b")
_DATE = re.compile(r"\b\d{4}-\d{2}-\d{2}\b")
FIXABLE = ("overflow", "off_slide", "dark_background")
BLOCKING = FIXABLE + ("table_overflow",)


def _slide_text(slide) -> str:
    parts = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            parts.append(sh.text_frame.text)
        if getattr(sh, "has_table", False) and sh.has_table:
            parts += [c.text for row in sh.table.rows for c in row.cells]
    return "\n".join(parts)


def _with_rounding(figures: set[str]) -> set[str]:
    """A fact like 60.4% may honestly appear as 60%; allow the half-up integer of every decimal figure."""
    out = set(figures)
    for f in figures:
        pct = f.endswith("%")
        num = f.rstrip("%")
        if "." in num:
            rounded = str(int(float(num) + 0.5))
            out.add(rounded + ("%" if pct else ""))
    return out


def fact_check(pptx_path: Path, facts: list[str], skip_slides: set[int]) -> list[dict]:
    allowed = _with_rounding(_figures(" ".join(facts)))
    issues = []
    for n, slide in enumerate(Presentation(str(pptx_path)).slides, start=1):
        if n in skip_slides:
            continue
        text = _slide_text(slide)
        cleaned = _DATE.sub(" ", _YEAR.sub(" ", _CITATION.sub(" ", text)))
        for figure in sorted(_figures(cleaned) - allowed):
            line = next((ln for ln in text.splitlines() if figure.rstrip("%") in ln), "")
            issues.append({"slide": n, "figure": figure, "text": line[:120]})
    return issues


def autofix(pptx_path: Path, issues: list[dict]) -> int:
    prs = Presentation(str(pptx_path))
    fixed = 0
    for issue in issues:
        slide = prs.slides[issue["slide"] - 1]
        if issue["kind"] == "overflow":
            for sh in slide.shapes:
                if sh.has_text_frame and sh.text_frame.text.startswith(issue["detail"]):
                    for p in sh.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(max(MIN_FONT_PT, (r.font.size.pt if r.font.size else 12) - SHRINK_PT))
                    fixed += 1
                    break
        elif issue["kind"] == "off_slide":
            for sh in slide.shapes:
                if sh.name == issue["detail"]:
                    sh.left = max(0, min(sh.left, prs.slide_width - sh.width))
                    sh.top = max(0, min(sh.top, prs.slide_height - sh.height))
                    fixed += 1
        elif issue["kind"] == "dark_background":
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            fixed += 1
    if fixed:
        prs.save(str(pptx_path))
    return fixed


def run_qc(pptx_path: Path, facts: list[str], skip_slides: set[int], png_dir: Path) -> dict:
    fixed = 0
    layout = qc.check_layout(pptx_path, skip=skip_slides)
    for _ in range(MAX_FIX_ROUNDS):
        fixable = [i for i in layout if i["kind"] in FIXABLE]
        if not fixable:
            break
        fixed += autofix(pptx_path, fixable)
        layout = qc.check_layout(pptx_path, skip=skip_slides)
    fact_issues = fact_check(pptx_path, facts, skip_slides)
    try:
        pngs = [str(p) for p in qc.export_pngs(pptx_path, png_dir)]
    except Exception as e:   # PowerPoint not installed / COM unavailable: thumbnails are optional
        logger.warning("slide PNG export unavailable: %s", e)
        pngs = []
    blocking = [i for i in layout if i["kind"] in BLOCKING]
    return {"layout": layout, "facts": fact_issues, "fixed": fixed, "pngs": pngs,
            "ready": not blocking and not fact_issues}
