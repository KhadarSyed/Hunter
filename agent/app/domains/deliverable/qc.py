"""Layout QC: geometry checks on text boxes + PowerPoint PNG export for visual review."""
from __future__ import annotations

import os
import subprocess
from pathlib import Path

from pptx import Presentation

EMU_PER_IN = 914400
CHAR_W_FACTOR = 0.5     # average Arial glyph width ≈ 0.5 × font size
LINE_H_FACTOR = 1.2
OVERLAP_EMU = 45720     # 0.05 in
EDGE_TOLERANCE_EMU = 9144
DARK_LUMINANCE = 0.6
TEXT_BOX = 17
SOLID_FILL = 1
EXPORT_TIMEOUT_S = 600


def estimate_overflow(text: str, width_in: float, height_in: float, font_pt: float) -> bool:
    chars_per_line = max(1, int((width_in * 72 - 14) / (font_pt * CHAR_W_FACTOR)))
    lines = sum(max(1, -(-len(p) // chars_per_line)) for p in text.split("\n"))
    return lines * font_pt * LINE_H_FACTOR > height_in * 72 - 7


def _luminance(hex_rgb: str) -> float:
    c = [int(hex_rgb[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    c = [x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4 for x in c]
    return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]


def _font_pt(shape) -> float:
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size:
                return r.font.size.pt
    return 12.0


def check_layout(pptx_path: Path, skip: set[int] | None = None) -> list[dict]:
    """skip: 1-based slide numbers to ignore (e.g. inherited template cover/closing slides)."""
    prs = Presentation(pptx_path)
    sw, sh_ = prs.slide_width, prs.slide_height
    issues = []
    for n, slide in enumerate(prs.slides, start=1):
        if skip and n in skip:
            continue
        fill = slide.background.fill
        if fill.type == SOLID_FILL and _luminance(str(fill.fore_color.rgb)) < DARK_LUMINANCE:
            issues.append({"slide": n, "kind": "dark_background", "detail": str(fill.fore_color.rgb)})
        for s in slide.shapes:
            if s.left is None:
                continue
            if (s.left < 0 or s.top < 0 or s.left + s.width > sw + EDGE_TOLERANCE_EMU
                    or s.top + s.height > sh_ + EDGE_TOLERANCE_EMU):
                issues.append({"slide": n, "kind": "off_slide", "detail": s.name})
        boxes = [s for s in slide.shapes if s.shape_type == TEXT_BOX and s.text_frame.text.strip()]
        for b in boxes:
            if estimate_overflow(b.text_frame.text, b.width / EMU_PER_IN, b.height / EMU_PER_IN, _font_pt(b)):
                issues.append({"slide": n, "kind": "overflow", "detail": b.text_frame.text[:50]})
        for i, a in enumerate(boxes):
            for b in boxes[i + 1:]:
                ox = min(a.left + a.width, b.left + b.width) - max(a.left, b.left)
                oy = min(a.top + a.height, b.top + b.height) - max(a.top, b.top)
                if ox > OVERLAP_EMU and oy > OVERLAP_EMU:
                    issues.append({"slide": n, "kind": "overlap",
                                   "detail": f"{a.text_frame.text[:25]!r} x {b.text_frame.text[:25]!r}"})
    return issues


def export_pngs(pptx_path: Path, out_dir: Path) -> list[Path]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    for old in out_dir.glob("slide_*.png"):
        old.unlink()
    # Paths travel as environment variables: PowerShell treats curly quotes (U+2018/2019, as in "Johnson’s")
    # as string delimiters, so interpolating them into the script breaks parsing and allows injection.
    env = {**os.environ, "QC_SRC": str(Path(pptx_path).resolve()), "QC_DST": str(out_dir.resolve())}
    script = (
        "$pp=New-Object -ComObject PowerPoint.Application;"
        "$p=$pp.Presentations.Open($env:QC_SRC,$true,$false,$false);"
        "$i=1;foreach($s in $p.Slides){"
        "$s.Export((Join-Path $env:QC_DST ('slide_' + $i.ToString('00') + '.png')),'PNG',1600,900);$i++};"
        "$p.Close();$pp.Quit()"
    )
    subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True, timeout=EXPORT_TIMEOUT_S, env=env)
    return sorted(out_dir.glob("slide_*.png"))
