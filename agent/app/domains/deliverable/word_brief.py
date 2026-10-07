"""The Word brief mirrors the deck: answers, per-RQ facts and cited insights, takeaways, method, citations."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Inches

from .generic_deck import DeckInput


def _table(doc, header: list[str], rows: list[list]):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for cell, h in zip(t.rows[0].cells, header):
        cell.text = h
    for r in rows:
        for cell, v in zip(t.add_row().cells, r):
            cell.text = str(v)
    return t


def _heading(doc, text: str, level: int, font: str):
    h = doc.add_heading(text, level)
    for run in h.runs:
        run.font.name = font
    return h


def _cited(card: dict) -> str:
    cites = "".join(f"[{c}]" for c in card.get("citations") or [])
    headline, text = card.get("headline", ""), card.get("text", "")
    return f"{headline}: {text} {cites}".strip()


def build_word_brief(inp: DeckInput, out_path: Path, slide_pngs: list[Path] | None = None) -> Path:
    doc = Document()
    if inp.brand_image and inp.brand_image.exists():
        doc.add_picture(str(inp.brand_image), width=Inches(6))
    _heading(doc, f"{inp.title} - {inp.subtitle}", 0, inp.title_font)
    doc.add_paragraph(f"{inp.date_label}  |  {inp.period_label}  |  Base: {inp.base_n} unique articles")
    _heading(doc, "Executive summary", 1, inp.title_font)
    _table(doc, ["Question", "Asked", "Answer"], [[a["rq_id"], a["question"], a["answer"]] for a in inp.answers])
    for i, rq in enumerate(inp.rqs):
        _heading(doc, f"{rq.id}: {rq.question}", 1, inp.title_font)
        sections = [s for s in inp.sections_by_rq.get(rq.id, []) if not s.skipped]
        if not sections:
            doc.add_paragraph("No articles for this question")
            continue
        for s in sections:
            _heading(doc, s.title, 2, inp.title_font)
            for fact in s.facts[:6]:
                doc.add_paragraph(fact, style="List Bullet")
        for card in inp.insights_by_rq.get(rq.id, []):
            doc.add_paragraph(_cited(card))
        png = slide_pngs[i] if slide_pngs and i < len(slide_pngs) else None
        if png and Path(png).exists():
            doc.add_picture(str(png), width=Inches(6))
    _heading(doc, "Key takeaways", 1, inp.title_font)
    for card in inp.takeaways:
        doc.add_paragraph(_cited(card), style="List Bullet")
    _heading(doc, "Methodology", 1, inp.title_font)
    for line in inp.methodology:
        doc.add_paragraph(line)
    _heading(doc, "Citations", 1, inp.title_font)
    _table(doc, ["#", "Outlet", "Headline", "Date", "URL"],
           [[c["n"], c["outlet"], c["title"], c["date"], c["url"]] for c in inp.citations])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
