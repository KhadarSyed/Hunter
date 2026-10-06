"""python-pptx building blocks in the Hunter style (light backgrounds only)."""
from __future__ import annotations

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from . import style

PERCENT_FORMAT = '0.0"%"'


def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def citation_suffix(cites: list[int]) -> str:
    return " " + "".join(f"[{c}]" for c in cites) if cites else ""


def new_content_slide(prs):
    layout = next(l for l in prs.slide_layouts if l.name == "Blank")
    slide = prs.slides.add_slide(layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(style.WHITE)
    return slide


def _rect(slide, x, y, w, h, fill: str, line: str | None = None, shape=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = _rgb(fill)
    if line:
        sh.line.color.rgb = _rgb(line)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def add_text(slide, x, y, w, h, text: str, size: float, bold=False, color=style.BODY,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size, r.font.bold = style.FONT, Pt(size), bold
        r.font.color.rgb = _rgb(color)
    return tb


def add_header(slide, kicker: str, title: str, summary: str):
    _rect(slide, 0, 0, style.SLIDE_W, style.HEADER_H, style.HEADER_BAND)
    add_text(slide, style.MARGIN, 0.15, 5.7, 0.35, kicker, style.KICKER_PT, True, style.VIOLET)
    add_text(slide, style.MARGIN, 0.5, 5.7, 0.8, title, style.TITLE_PT, True, "1F1F1F")
    add_text(slide, 6.35, 0.18, 6.55, 1.05, summary, style.SUMMARY_PT, True, style.BODY, anchor=MSO_ANCHOR.MIDDLE)


def add_footer(slide, source: str, base_n: int):
    _rect(slide, 0, style.FOOTER_Y, style.SLIDE_W, style.FOOTER_H, style.FOOTER_BAND)
    add_text(slide, 0.27, style.FOOTER_Y + 0.01, 6.5, 0.22, source.upper(), style.FOOTER_PT, False, style.FOOTER_TEXT)
    add_text(slide, 7.0, style.FOOTER_Y + 0.01, 6.1, 0.22,
             f"Base: N={base_n} unique articles collected (Meltwater + manual extraction)",
             style.FOOTER_PT, False, style.FOOTER_TEXT, align=PP_ALIGN.RIGHT)


def _chart(slide, kind, x, y, w, h, categories, values, number_format):
    data = CategoryChartData(number_format=number_format)
    data.categories = categories
    data.add_series("Articles", values)
    chart = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    chart.has_legend = False
    chart.has_title = False
    chart.font.name, chart.font.size = style.FONT, Pt(9)
    chart.font.color.rgb = _rgb(style.BODY)
    return chart


def _style_axes(chart):
    chart.value_axis.visible = False
    chart.value_axis.has_major_gridlines = False
    chart.category_axis.format.line.color.rgb = _rgb(style.CARD_LINE)
    chart.category_axis.tick_labels.font.size = Pt(9)


def _labels(plot, number_format):
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.show_value = True
    dl.number_format, dl.number_format_is_linked = number_format, False
    dl.font.size, dl.font.name = Pt(9), style.FONT
    dl.font.color.rgb = _rgb(style.BODY)
    return dl


def add_column_chart(slide, x, y, w, h, categories, values, highlight_idx, number_format="0"):
    chart = _chart(slide, XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, w, h, categories, values, number_format)
    _style_axes(chart)
    plot = chart.plots[0]
    plot.gap_width = 60
    _labels(plot, number_format).position = XL_LABEL_POSITION.OUTSIDE_END
    for i, pt in enumerate(plot.series[0].points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = _rgb(style.VIOLET if i in highlight_idx else style.LAVENDER)
    return chart


def add_bar_chart(slide, x, y, w, h, categories, values, color=style.LAVENDER, number_format="0"):
    pairs = list(zip(categories, values))[::-1]          # PowerPoint draws bars bottom-up; largest on top
    chart = _chart(slide, XL_CHART_TYPE.BAR_CLUSTERED, x, y, w, h, [c for c, _ in pairs], [v for _, v in pairs],
                   number_format)
    _style_axes(chart)
    plot = chart.plots[0]
    plot.gap_width = 50
    _labels(plot, number_format).position = XL_LABEL_POSITION.OUTSIDE_END
    plot.series[0].format.fill.solid()
    plot.series[0].format.fill.fore_color.rgb = _rgb(color)
    return chart


def add_line_chart_with_peaks(slide, x, y, w, h, categories, values, peak_idx):
    chart = _chart(slide, XL_CHART_TYPE.LINE_MARKERS, x, y, w, h, categories, values, "0")
    _style_axes(chart)
    series = chart.plots[0].series[0]
    series.smooth = False      # smoothing overshoots below zero between low months
    series.format.line.color.rgb = _rgb(style.POWDER_BLUE)
    series.format.line.width = Pt(2)
    series.marker.style = XL_MARKER_STYLE.NONE
    for i in peak_idx:
        pt = series.points[i]
        pt.marker.style, pt.marker.size = XL_MARKER_STYLE.CIRCLE, 9
        pt.marker.format.fill.solid()
        pt.marker.format.fill.fore_color.rgb = _rgb(style.VIOLET)
        pt.marker.format.line.color.rgb = _rgb(style.VIOLET)
        dl = pt.data_label
        dl.position = XL_LABEL_POSITION.ABOVE      # python-pptx's point dLbl defaults showVal=1
        dl.font.bold, dl.font.size = True, Pt(10)
        dl.font.color.rgb = _rgb(style.VIOLET)
    return chart


def add_doughnut(slide, x, y, w, h, categories, values, number_format=PERCENT_FORMAT):
    chart = _chart(slide, XL_CHART_TYPE.DOUGHNUT, x, y, w, h, categories, values, number_format)
    chart.has_legend = True
    chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.RIGHT, False
    chart.legend.font.size = Pt(9)
    plot = chart.plots[0]
    _labels(plot, number_format)
    for i, pt in enumerate(plot.series[0].points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = _rgb(style.PALETTE[i % len(style.PALETTE)])
    return chart


def _card(slide, x, y, w, h):
    card = _rect(slide, x, y, w, h, style.CARD_FILL, style.CARD_LINE, MSO_SHAPE.ROUNDED_RECTANGLE)
    card.adjustments[0] = 0.06
    return card


def _headline_height(text: str, width_in: float) -> float:
    chars_per_line = max(1, int((width_in * 72 - 14) / (style.CARD_HEAD_PT * 0.5)))
    lines = max(1, -(-len(text) // chars_per_line))
    return style.CARD_HEAD_H if lines > 1 else style.CARD_HEAD_H / 2 + 0.05


def add_insight_cards(slide, x, y, w, h, insights: list[dict], cols: int):
    if not insights:
        return
    gap = 0.15
    rows = -(-len(insights) // cols)
    cw, ch = (w - gap * (cols - 1)) / cols, (h - gap * (rows - 1)) / rows
    for i, ins in enumerate(insights):
        cx, cy = x + (i % cols) * (cw + gap), y + (i // cols) * (ch + gap)
        _card(slide, cx, cy, cw, ch)
        head_h = _headline_height(ins["headline"], cw - 0.3)
        add_text(slide, cx + 0.15, cy + 0.08, cw - 0.3, head_h, ins["headline"], style.CARD_HEAD_PT, True, style.VIOLET)
        body_top = 0.08 + head_h + 0.04
        add_text(slide, cx + 0.15, cy + body_top, cw - 0.3, ch - body_top - 0.08,
                 ins["text"] + citation_suffix(ins["citations"]), style.CARD_BODY_PT)


def add_kpi_tiles(slide, x, y, w, h, tiles: list[dict]):
    gap = 0.2
    tw = (w - gap * (len(tiles) - 1)) / len(tiles)
    for i, t in enumerate(tiles):
        tx = x + i * (tw + gap)
        _card(slide, tx, y, tw, h)
        add_text(slide, tx + 0.15, y + 0.12, tw - 0.3, 0.8, t["value"], 32, True, style.VIOLET, PP_ALIGN.CENTER)
        add_text(slide, tx + 0.15, y + 0.95, tw - 0.3, 0.4, t["label"], 11, True, "1F1F1F", PP_ALIGN.CENTER)
        add_text(slide, tx + 0.15, y + 1.38, tw - 0.3, h - 1.48, t["note"], style.BODY_PT, False, style.BODY,
                 PP_ALIGN.CENTER)


def add_table(slide, x, y, w, h, header: list[str], rows: list[list[str]], col_widths: list[float]):
    table = slide.shapes.add_table(len(rows) + 1, len(header), Inches(x), Inches(y), Inches(w), Inches(h)).table
    for j, cw in enumerate(col_widths):
        table.columns[j].width = Inches(cw)
    for r, values in enumerate([header] + rows):
        for c, v in enumerate(values):
            cell = table.cell(r, c)
            cell.text = str(v)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(style.FOOTER_BAND if r == 0 else (style.WHITE if r % 2 else style.CARD_FILL))
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.name, run.font.size, run.font.bold = style.FONT, Pt(9 if r == 0 else 8), r == 0
                    run.font.color.rgb = _rgb(style.VIOLET if r == 0 else style.BODY)
    return table
