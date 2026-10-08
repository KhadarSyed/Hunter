"""Deck validator: after the deck is built, every requirement of a client deck is checked and each gap named --
the client's logo and product on the cover, the brand vs competitor chapter, an answer, a chart, a cited summary
and verbatims for every question, a cited summary on every chart slide, citations that name a brand literally (not
an idiom like "a band-aid solution") and are listed (with their icon) in Sources, competitors' verbatims on their own
slide with the post or article screenshot, icons beside chart labels, the social vs editorial split, slides that
build on click in PowerPoint, and the brief's own asks.

Rows are {check, status: pass | warn | fail, detail}; a fail is something the client would notice is missing.
"""
from __future__ import annotations

from pathlib import Path

from ..deliverable.engine_types import RQ
from .spec import DeckSpec, SlideSpec


def _row(check: str, ok: bool, detail: str = "", soft: bool = False) -> dict:
    return {"check": check, "status": "pass" if ok else ("warn" if soft else "fail"), "detail": "" if ok else detail}


def _cited(slides: list[SlideSpec]) -> set[str]:
    return {str(n) for s in slides for c in s.cards for n in c.get("citations") or []}


def _question(rq: RQ, slides: list[SlideSpec]) -> dict:
    mine = [s for s in slides if s.id.startswith(rq.id.lower() + "-")]
    divider = next((s for s in mine if s.type == "divider"), None)
    gaps = []
    if not divider or not (divider.so_what or "").strip():
        gaps.append("no answer")
    if not any(s.charts or s.tables for s in mine if s.type != "verbatim_wall"):
        gaps.append("no chart")
    if not _cited([s for s in mine if s.type != "verbatim_wall"]):
        gaps.append("no cited summary")
    if not any(s.type == "verbatim_wall" and s.cards for s in mine):
        gaps.append("no verbatims")
    return _row(rq.question, not gaps, ", ".join(gaps))


def _chart_slides(slides: list[SlideSpec], rqs: list[RQ]) -> list[SlideSpec]:
    """The evidence slides: a chart or table in the brand chapter or a question's chapter."""
    prefixes = ("brand-",) + tuple(rq.id.lower() + "-" for rq in rqs)
    return [s for s in slides if s.type not in ("verbatim_wall", "divider") and (s.charts or s.tables)
            and s.id.startswith(prefixes)]


def _evidence_rows(slides: list[SlideSpec], rqs: list[RQ], has_competitors: bool,
                   citation_brands: dict[str, list[str]] | None) -> list[dict]:
    charts = _chart_slides(slides, rqs)
    bare = [s.id for s in charts if not _cited([s])]
    rows = [_row("Every chart slide has a cited summary", not bare, "no cited summary on " + ", ".join(bare[:6]))]
    unlabelled = [s.id for s in charts
                  if (cats := [c for ch in s.charts for c in ch.get("categories") or []]) and not set(cats) & set(s.logos)]
    rows.append(_row("Chart labels show their icon or logo", not unlabelled,
                     f"{len(unlabelled)} chart(s) without icons: " + ", ".join(unlabelled[:4]), soft=True))
    if citation_brands is not None:
        stray = sorted((n for n in _cited(slides) if not citation_brands.get(n)), key=lambda n: int(n) if n.isdigit() else 0)
        rows.append(_row("Citations name the brand or a competitor literally", not stray,
                         "cited source(s) never name a brand: " + ", ".join(stray[:12])))
    walls = [s for s in slides if s.type == "verbatim_wall"]
    if has_competitors:
        rows.append(_row("Competitor verbatims on their own slide",
                         any(s.id.endswith("competitors-verbatims") and s.cards for s in walls),
                         "no competitor verbatim slide", soft=True))
    no_shot = sum(1 for s in walls for c in s.cards if not c.get("image"))
    rows.append(_row("Verbatims show the post or article screenshot", not no_shot,
                     f"{no_shot} verbatim(s) without a screenshot", soft=True))
    rows.append(_row("Social vs editorial split", any("social vs editorial" in (s.title or "").lower() for s in charts),
                     "no social vs editorial chart", soft=True))
    return rows


def _builds_row(spec: DeckSpec, pptx: Path) -> dict:
    """Slides with cards or verbatims must reveal them on click (a <p:timing> in the slide XML)."""
    from pptx import Presentation
    ns = "{http://schemas.openxmlformats.org/presentationml/2006/main}timing"
    try:
        animated = [s._element.find(ns) is not None for s in Presentation(str(pptx)).slides]
    except Exception as e:                  # an unreadable file is itself a failure to report, not to raise
        return _row("Slides build on click in PowerPoint", False, f"could not read the PPTX: {type(e).__name__}")
    wanted = [k for k, s in enumerate(spec.slides) if s.cards and k < len(animated)]
    still = [spec.slides[k].id for k in wanted if not animated[k]]
    if wanted and len(still) == len(wanted):
        return _row("Slides build on click in PowerPoint", False, "no slide animates its cards")
    return _row("Slides build on click in PowerPoint", not still, f"{len(still)} slide(s) without builds: "
                + ", ".join(still[:4]), soft=True)


def validate_deck(spec: DeckSpec, rqs: list[RQ], checklist: list[dict], brand_products: list[str],
                  has_competitors: bool, citation_brands: dict[str, list[str]] | None = None,
                  pptx: Path | None = None, is_category: bool = False) -> list[dict]:
    """`is_category`: the client is a category, not a brand -- no single logo, and coverage needn't name a brand."""
    slides = spec.slides
    rows = [_row("Client logo on the cover", bool(spec.brand_logo),
                 "a category study has no single brand logo" if is_category else "no logo for the client brand",
                 soft=is_category),
            _row("Client product on the cover", bool(spec.product_image) or not brand_products,
                 f"no photo found for {', '.join(brand_products[:2])}", soft=True)]
    if has_competitors:
        rows.append(_row("Brand share of voice vs competitors", any(s.id.startswith("brand-sov") for s in slides),
                         "no share of voice slide"))
        rows.append(_row("Brand sentiment vs competitors", any(s.id.startswith("brand-sentiment") for s in slides),
                         "no sentiment slide"))
    rows += [_question(rq, slides) for rq in rqs]
    rows += _evidence_rows(slides, rqs, has_competitors, None if is_category else citation_brands)
    summary = next((s for s in slides if s.type == "executive_summary"), None)
    uncited = [c.get("headline") or "" for c in (summary.cards if summary else []) if not c.get("citations")]
    rows.append(_row("Executive summary cites its sources", bool(summary) and not uncited,
                     f"{len(uncited)} answer(s) without a source" if summary else "no executive summary", soft=True))
    listed = {str(r[0]) for s in slides if s.type == "citations" for t in s.tables for r in t.get("rows", []) if r}
    missing = sorted(_cited(slides) - listed, key=lambda n: int(n) if n.isdigit() else 0)
    rows.append(_row("Every citation is listed in Sources", not missing, "not in Sources: " + ", ".join(missing)))
    no_icon = sorted(n for n in _cited(slides) if not (spec.citation_meta.get(n) or {}).get("icon"))
    rows.append(_row("Citations show their source icon", not no_icon, f"{len(no_icon)} source(s) without an icon",
                     soft=True))
    if pptx:
        rows.append(_builds_row(spec, pptx))
    rows += [_row(f"Brief: {r['ask']}", r["status"] == "covered", r.get("note") or r["status"],
                  soft=r["status"] == "partial")
             for r in checklist if r.get("kind") != "question"]
    return rows
