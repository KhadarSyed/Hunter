"""Pixel-perfect exports of the HTML deck: each slide screenshotted at 1920x1080 into a 16:9 PPTX and a PDF.

In the PPTX a slide is presented the way a presenter would build it: the slide picture without its cards, then
each card (insight, verbatim) as its own picture at its exact place, fading in on click; slides cross-fade.
The title stays under the pictures for search and accessibility, and the slide text goes in the notes. The PDF
keeps every slide whole.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Inches

from .spec import DeckSpec

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
PX_W, PX_H = 1920, 1080
SETTLE_MS = 900
MAX_BUILDS = 9
FADE_MS = 500
BUILD_SELECTOR = ".build"
_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"


@dataclass
class Shot:
    full: Path
    base: Path | None = None                          # the slide with its builds hidden
    builds: list[tuple[Path, tuple[float, float, float, float]]] = field(default_factory=list)   # (png, x, y, w, h px)


def _capture(el, slide_dir: Path, k: int) -> Shot:
    shot = Shot(full=slide_dir / f"slide_{k:02d}.png")
    el.screenshot(path=str(shot.full))
    origin = el.bounding_box()
    builds = el.query_selector_all(BUILD_SELECTOR)[:MAX_BUILDS]
    if not builds or not origin:
        return shot
    for i, b in enumerate(builds, start=1):
        box = b.bounding_box()
        if not box or box["width"] < 4 or box["height"] < 4:
            continue
        path = slide_dir / "builds" / f"slide_{k:02d}_build_{i}.png"
        path.parent.mkdir(parents=True, exist_ok=True)
        b.screenshot(path=str(path))
        shot.builds.append((path, (box["x"] - origin["x"], box["y"] - origin["y"], box["width"], box["height"])))
    if shot.builds:
        el.evaluate(f"s => s.querySelectorAll('{BUILD_SELECTOR}').forEach(b => b.style.visibility = 'hidden')")
        shot.base = slide_dir / "builds" / f"slide_{k:02d}_base.png"
        el.screenshot(path=str(shot.base))
        el.evaluate(f"s => s.querySelectorAll('{BUILD_SELECTOR}').forEach(b => b.style.visibility = '')")
    return shot


def screenshot_slides(html_path: Path, out_dir: Path) -> list[Shot]:
    from playwright.sync_api import sync_playwright
    out_dir.mkdir(parents=True, exist_ok=True)
    shots = []
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": PX_W, "height": PX_H})
        page.goto(Path(html_path).resolve().as_uri() + "?export=1", wait_until="networkidle")
        page.wait_for_timeout(SETTLE_MS)          # web fonts and photos settle
        for k, el in enumerate(page.query_selector_all("section.slide"), start=1):
            shots.append(_capture(el, out_dir, k))
        browser.close()
    return shots


def _emu(px: float, total_px: int, total_emu: int) -> Emu:
    return Emu(int(round(px / total_px * total_emu)))


def _p(tag: str, **attrs) -> etree._Element:
    el = etree.Element(f"{{{_NS}}}{tag}")
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def _sub(parent, tag: str, **attrs) -> etree._Element:
    el = _p(tag, **attrs)
    parent.append(el)
    return el


def _click_fades(shape_ids: list[int]) -> etree._Element:
    """<p:timing>: one click per shape, each a fade-in entrance (PowerPoint's preset 10)."""
    ids = iter(range(1, 10_000))
    timing = _p("timing")
    root = _sub(_sub(_sub(timing, "tnLst"), "par"), "cTn", id=next(ids), dur="indefinite", restart="never",
                nodeType="tmRoot")
    seq = _sub(_sub(root, "childTnLst"), "seq", concurrent="1", nextAc="seek")
    main = _sub(seq, "cTn", id=next(ids), dur="indefinite", nodeType="mainSeq")
    clicks = _sub(main, "childTnLst")
    for spid in shape_ids:
        click = _sub(_sub(clicks, "par"), "cTn", id=next(ids), fill="hold")
        _sub(_sub(click, "stCondLst"), "cond", delay="indefinite")
        step = _sub(_sub(_sub(click, "childTnLst"), "par"), "cTn", id=next(ids), fill="hold")
        _sub(_sub(step, "stCondLst"), "cond", delay="0")
        effect = _sub(_sub(_sub(step, "childTnLst"), "par"), "cTn", id=next(ids), presetID="10", presetClass="entr",
                      presetSubtype="0", fill="hold", grpId="0", nodeType="clickEffect")
        _sub(_sub(effect, "stCondLst"), "cond", delay="0")
        behaviours = _sub(effect, "childTnLst")
        show = _sub(behaviours, "set")
        bhvr = _sub(show, "cBhvr")
        flip = _sub(bhvr, "cTn", id=next(ids), dur="1", fill="hold")
        _sub(_sub(flip, "stCondLst"), "cond", delay="0")
        _sub(_sub(bhvr, "tgtEl"), "spTgt", spid=spid)
        _sub(_sub(bhvr, "attrNameLst"), "attrName").text = "style.visibility"
        _sub(_sub(show, "to"), "strVal", val="visible")
        fade = _sub(behaviours, "animEffect", transition="in", filter="fade")
        fbhvr = _sub(fade, "cBhvr")
        _sub(fbhvr, "cTn", id=next(ids), dur=FADE_MS)
        _sub(_sub(fbhvr, "tgtEl"), "spTgt", spid=spid)
    prev = _sub(_sub(seq, "prevCondLst"), "cond", evt="onPrev", delay="0")
    _sub(_sub(prev, "tgtEl"), "sldTgt")
    nxt = _sub(_sub(seq, "nextCondLst"), "cond", evt="onNext", delay="0")
    _sub(_sub(nxt, "tgtEl"), "sldTgt")
    return timing


def _animate(slide, shape_ids: list[int]) -> None:
    """A fade transition into the slide, then the build pictures fading in one click at a time."""
    sld = slide._element
    for old in sld.findall(f"{{{_NS}}}transition") + sld.findall(f"{{{_NS}}}timing"):
        sld.remove(old)
    transition = _p("transition", spd="med")
    _sub(transition, "fade")
    ext = sld.find(f"{{{_NS}}}extLst")
    at = list(sld).index(ext) if ext is not None else len(sld)
    sld.insert(at, transition)
    if shape_ids:
        sld.insert(at + 1, _click_fades(shape_ids))


def build_pptx(pngs: list[Path], spec: DeckSpec, out: Path, builds: list[Shot] | None = None) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    layout = prs.slide_layouts[5]          # "Title Only": the title stays searchable under the picture
    for k, (png, s) in enumerate(zip(pngs, spec.slides)):
        shot = builds[k] if builds and k < len(builds) else None
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = s.question or s.title or spec.title
        slide.shapes.add_picture(str(shot.base if shot and shot.base else png), Emu(0), Emu(0), SLIDE_W, SLIDE_H)
        ids = []
        for path, (x, y, w, h) in (shot.builds if shot and shot.base else []):
            pic = slide.shapes.add_picture(str(path), _emu(x, PX_W, SLIDE_W), _emu(y, PX_H, SLIDE_H),
                                           _emu(w, PX_W, SLIDE_W), _emu(h, PX_H, SLIDE_H))
            ids.append(pic.shape_id)
        _animate(slide, ids)
        lines = [s.kicker, s.question or s.title, s.so_what] + [f"{c.get('headline', '')}: {c.get('text', '')}" for c in s.cards]
        if s.type in ("verbatim_wall", "citations") and s.notes:
            lines.append(s.notes)
        slide.notes_slide.notes_text_frame.text = "\n".join(x for x in lines if x)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


def build_pdf(pngs: list[Path], out: Path) -> Path:
    images = [Image.open(p).convert("RGB") for p in pngs]
    out.parent.mkdir(parents=True, exist_ok=True)
    images[0].save(out, "PDF", save_all=True, append_images=images[1:], resolution=144)
    return out


def export_all(html_path: Path, spec: DeckSpec, out_dir: Path, stem: str) -> dict:
    shots = screenshot_slides(html_path, out_dir / "slides")
    if len(shots) != len(spec.slides):      # never pair titles and notes with the wrong pictures
        raise RuntimeError(f"rendered {len(shots)} slides for a {len(spec.slides)}-slide deck")
    pngs = [s.full for s in shots]
    safe = "".join(ch for ch in stem if ch.isalnum() or ch in " -_").strip() or "Deck"
    return {"pngs": pngs, "pptx": build_pptx(pngs, spec, out_dir / f"{safe} - Deck.pptx", shots),
            "pdf": build_pdf(pngs, out_dir / f"{safe} - Deck.pdf")}
