"""Secondary research module: searches the web for recent news, trends, and
developments about the client/brand using context-aware queries, then produces
an analytical Word document with credible-source prioritization."""
from __future__ import annotations

import re
import time
import warnings
from datetime import date
from pathlib import Path
from typing import Callable, Optional

warnings.filterwarnings("ignore", message=".*renamed to.*ddgs.*")
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

EventFn = Callable[[str, dict], None]

VIOLET = RGBColor(0x5B, 0x2C, 0x9D)
GREY = RGBColor(0x6E, 0x6E, 0x6E)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
BLUE = RGBColor(0x00, 0x56, 0xB3)

CREDIBLE_DOMAINS = {
    "reuters.com", "apnews.com", "bloomberg.com", "cnbc.com",
    "wsj.com", "ft.com", "nytimes.com", "washingtonpost.com",
    "bbc.com", "bbc.co.uk", "theguardian.com",
    "forbes.com", "fortune.com", "businessinsider.com",
    "techcrunch.com", "wired.com", "theverge.com",
    "gartner.com", "forrester.com", "mckinsey.com",
    "deloitte.com", "pwc.com", "ey.com", "kpmg.com",
    "statista.com", "emarketer.com", "nielseniq.com",
    "sec.gov", "ftc.gov", "fda.gov", "epa.gov",
    "hbr.org", "economist.com", "marketwatch.com",
    "morningstar.com", "seekingalpha.com", "investopedia.com",
    "adage.com", "adweek.com", "prweek.com", "digiday.com",
}

SECTION_NEWS = "Latest News & Developments"
SECTION_INDUSTRY = "Industry Trends & Market Dynamics"
SECTION_COMPETITIVE = "Competitive Developments"
SECTION_CONSUMER = "Consumer Trends & Behavior"
SECTION_REGULATORY = "Regulatory & Policy Updates"
SECTION_MARKET_DATA = "Market Data & Forecasts"

SECTION_ORDER = [
    SECTION_NEWS,
    SECTION_INDUSTRY,
    SECTION_COMPETITIVE,
    SECTION_CONSUMER,
    SECTION_REGULATORY,
    SECTION_MARKET_DATA,
]


REGION_MAP = {
    "US": "us-en", "UK": "uk-en", "Global": "wt-wt",
    "North America": "us-en", "Europe": "eu-en",
    "APAC": "wt-wt", "LATAM": "wt-wt", "EMEA": "wt-wt",
}

BLOCKED_DOMAINS = {
    "wikipedia.org", "wikimedia.org", "wiktionary.org", "britannica.com",
    "flipkart.com", "amazon.in", "myntra.com", "snapdeal.com",
    "naukri.com", "indiamart.com", "justdial.com", "zomato.com",
    "swiggy.com", "bigbasket.com", "jiomart.com", "meesho.com",
    "facebook.com", "instagram.com", "twitter.com", "x.com",
    "youtube.com", "tiktok.com", "pinterest.com",
    "linkedin.com", "reddit.com", "quora.com",
    "etsy.com", "ebay.com", "aliexpress.com",
}

BLOCKED_TLDS = {".in", ".co.in", ".org.in", ".net.in", ".ind.in"}


def _search_web(query: str, max_results: int = 8, region: str = "us-en") -> list[dict]:
    """Run a DuckDuckGo web search and return results."""
    try:
        from ddgs import DDGS
        results = DDGS().text(query, max_results=max_results)
        return results
    except Exception:
        try:
            from duckduckgo_search import DDGS
            with DDGS() as ddgs:
                return list(ddgs.text(query, region=region, max_results=max_results))
        except Exception:
            return []


def _search_news(query: str, max_results: int = 8, region: str = "us-en") -> list[dict]:
    """Run a DuckDuckGo news search for recent articles."""
    try:
        from ddgs import DDGS
        results = DDGS().news(query, max_results=max_results)
        return results
    except Exception:
        try:
            from duckduckgo_search import DDGS
            with DDGS() as ddgs:
                return list(ddgs.news(query, region=region, max_results=max_results))
        except Exception:
            return []


def _is_blocked(url: str, client: str = "", geography: str = "US") -> bool:
    """Check if a URL should be excluded."""
    domain = _extract_domain(url).lower()
    if any(domain == bd or domain.endswith("." + bd) for bd in BLOCKED_DOMAINS):
        return True
    if geography in ("US", "Global", "North America", "Europe", "UK"):
        if any(domain.endswith(tld) for tld in BLOCKED_TLDS):
            return True
    if client:
        client_slug = re.sub(r"[^a-z0-9]", "", client.lower())
        if client_slug and client_slug in domain:
            return True
    return False


def _deduplicate(results: list[dict], client: str = "", geography: str = "US") -> list[dict]:
    """Remove duplicate URLs and blocked domains."""
    seen = set()
    deduped = []
    for r in results:
        url = r.get("href") or r.get("url") or r.get("link", "")
        if not url or url in seen:
            continue
        if _is_blocked(url, client, geography):
            continue
        seen.add(url)
        deduped.append(r)
    return deduped


def _extract_domain(url: str) -> str:
    """Extract domain name from URL for source attribution."""
    match = re.search(r"https?://(?:www\.)?([^/]+)", url)
    return match.group(1) if match else ""


def _is_credible(url: str) -> bool:
    """Check if a URL belongs to a credible source."""
    domain = _extract_domain(url).lower()
    return any(domain == cd or domain.endswith("." + cd) for cd in CREDIBLE_DOMAINS)


def _sort_by_credibility(results: list[dict]) -> list[dict]:
    """Sort results with credible sources first, preserving order within groups."""
    credible = []
    other = []
    for r in results:
        url = r.get("href") or r.get("url") or r.get("link", "")
        if _is_credible(url):
            credible.append(r)
        else:
            other.append(r)
    return credible + other


def _build_queries(
    client: str,
    competitors: list[str],
    category: str,
    topics: list[str],
    geography: str,
    time_period: str,
) -> list[dict]:
    """Construct context-aware search queries tagged with their intended section.

    Returns a list of dicts: {"query": str, "section": str, "use_news": bool}
    Queries are kept short and simple for better DuckDuckGo results.
    """
    year = date.today().year
    geo = geography or "US"
    comp_names = competitors[:3] if competitors else []

    queries = []

    def add(section: str, query: str, use_news: bool = False):
        queries.append({"query": query, "section": section, "use_news": use_news})

    # --- Latest News ---
    add(SECTION_NEWS, f"{client} news {year}", use_news=True)
    add(SECTION_NEWS, f"{client} earnings revenue {year}", use_news=True)
    add(SECTION_NEWS, f"{client} strategy announcement", use_news=True)

    # --- Industry Trends ---
    add(SECTION_INDUSTRY, f"{category} industry trends {year}", use_news=True)
    add(SECTION_INDUSTRY, f"{category} market outlook {year}")
    add(SECTION_INDUSTRY, f"{category} growth challenges {geo}")

    # --- Competitive ---
    if comp_names:
        add(SECTION_COMPETITIVE, f"{client} vs {comp_names[0]} market share", use_news=True)
        for comp in comp_names[:2]:
            add(SECTION_COMPETITIVE, f"{comp} news {year}", use_news=True)
    else:
        add(SECTION_COMPETITIVE, f"{category} top companies market share {year}")
        add(SECTION_COMPETITIVE, f"{client} competitors {category}")

    # --- Consumer ---
    add(SECTION_CONSUMER, f"{category} consumer trends {year}")
    add(SECTION_CONSUMER, f"{client} customer experience {year}", use_news=True)

    # --- Regulatory ---
    add(SECTION_REGULATORY, f"{category} regulation policy {geo} {year}")
    add(SECTION_REGULATORY, f"{category} FDA compliance {year}" if "food" in category.lower() or "beverage" in category.lower()
        else f"{category} government regulation {year}")

    # --- Market Data ---
    add(SECTION_MARKET_DATA, f"{category} market size {year}")
    add(SECTION_MARKET_DATA, f"{category} industry forecast revenue growth")

    # --- Topic-specific ---
    for topic in topics[:3]:
        add(SECTION_INDUSTRY, f"{client} {topic.lower()} {year}", use_news=True)

    return queries


def run_research(
    client: str,
    competitors: list[str],
    category: str,
    topics: list[str],
    output_path: str,
    on_event: Optional[EventFn] = None,
    geography: str = "US",
    time_period: str = "",
) -> dict:
    """Execute secondary research and produce a Word document.

    Returns dict with status, output_path, source_count, and search_queries.
    """
    emit = on_event or (lambda et, p: None)
    all_results: list[dict] = []
    queries_run: list[str] = []

    emit("research_started", {"client": client})

    region = REGION_MAP.get(geography, "us-en")
    tagged_queries = _build_queries(client, competitors, category, topics, geography, time_period)

    for i, tq in enumerate(tagged_queries):
        query = tq["query"]
        section = tq["section"]
        use_news = tq.get("use_news", False)
        emit("research_searching", {"query": query, "index": i + 1, "total": len(tagged_queries)})

        web_results = _search_web(query, max_results=8, region=region)
        time.sleep(2)

        if use_news:
            news_results = _search_news(query, max_results=6, region=region)
            time.sleep(2)
        else:
            news_results = []

        for r in web_results:
            r["_query"] = query
            r["_section"] = section
            r["_type"] = "web"
        for r in news_results:
            r["_query"] = query
            r["_section"] = section
            r["_type"] = "news"

        all_results.extend(news_results)
        all_results.extend(web_results)
        queries_run.append(query)

    all_results = _deduplicate(all_results, client=client, geography=geography)
    emit("research_results", {"total_sources": len(all_results)})

    grouped = _group_results(all_results, client, competitors, category)

    emit("research_writing", {"output_path": output_path})
    _build_document(
        client=client,
        competitors=competitors,
        category=category,
        geography=geography,
        grouped=grouped,
        all_results=all_results,
        queries_run=queries_run,
        output_path=output_path,
    )

    emit("research_complete", {"output_path": output_path, "source_count": len(all_results)})

    return {
        "status": "completed",
        "output_path": output_path,
        "source_count": len(all_results),
        "search_queries": queries_run,
    }


def _group_results(
    results: list[dict],
    client: str,
    competitors: list[str],
    category: str,
) -> dict[str, list[dict]]:
    """Organize results into thematic sections using query tags and content analysis."""
    groups: dict[str, list[dict]] = {s: [] for s in SECTION_ORDER}

    client_lower = client.lower()
    comp_lower = [c.lower() for c in competitors]

    for r in results:
        tagged_section = r.get("_section", "")

        title = (r.get("title") or "").lower()
        body = (r.get("body") or r.get("snippet") or r.get("description") or "").lower()
        text = f"{title} {body}"

        section = _classify_result(tagged_section, text, client_lower, comp_lower)
        groups[section].append(r)

    max_per_section = 10
    for section in groups:
        groups[section] = _sort_by_credibility(groups[section])[:max_per_section]

    return {k: v for k, v in groups.items() if v}


def _classify_result(
    tagged_section: str,
    text: str,
    client_lower: str,
    comp_lower: list[str],
) -> str:
    """Determine the best section for a result using tag + content signals.

    Trusts the query tag by default. Only overrides when content has 2+ strong
    signals pointing to a different section.
    """
    section_signals = {
        SECTION_REGULATORY: ["regulat", "compliance", "fda", "ftc", "legislation", "mandate", "ruling"],
        SECTION_MARKET_DATA: ["market size", "forecast", "billion", "cagr", "valuation", "projection"],
        SECTION_COMPETITIVE: [" vs ", "versus", "market share", "rival"],
        SECTION_CONSUMER: ["consumer trend", "customer experience", "satisfaction survey", "buyer behavior"],
    }

    if tagged_section and tagged_section in SECTION_ORDER:
        best = tagged_section
        best_count = 0
        for section, signals in section_signals.items():
            if section == tagged_section:
                continue
            hits = sum(1 for s in signals if s in text)
            if hits >= 2 and hits > best_count:
                best = section
                best_count = hits
        return best

    for section, signals in section_signals.items():
        if sum(1 for s in signals if s in text) >= 2:
            return section

    if any(c in text for c in comp_lower if len(c) > 3):
        return SECTION_COMPETITIVE

    return SECTION_NEWS


# ---------------------------------------------------------------------------
# Document building
# ---------------------------------------------------------------------------

def _build_document(
    client: str,
    competitors: list[str],
    category: str,
    geography: str,
    grouped: dict[str, list[dict]],
    all_results: list[dict],
    queries_run: list[str],
    output_path: str,
) -> None:
    """Build the research Word document."""
    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)
    font.color.rgb = BLACK

    for heading_level in range(1, 4):
        h_style = doc.styles[f"Heading {heading_level}"]
        h_font = h_style.font
        h_font.name = "Calibri"
        h_font.color.rgb = VIOLET

    _add_cover_page(doc, client, category)
    doc.add_page_break()

    _add_executive_summary(doc, client, competitors, category, geography, grouped, all_results)
    doc.add_page_break()

    for section_title in SECTION_ORDER:
        if section_title in grouped:
            _add_findings_section(doc, section_title, grouped[section_title], client)

    doc.add_page_break()
    _add_opportunities_risks(doc, client, category, grouped)

    doc.add_page_break()
    _add_source_appendix(doc, all_results, queries_run)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def _add_cover_page(doc: Document, client: str, category: str) -> None:
    """Add a branded cover page."""
    for _ in range(6):
        doc.add_paragraph("")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("SECONDARY RESEARCH REPORT")
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = VIOLET

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(client)
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = VIOLET

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{category} | {date.today().strftime('%B %Y')}")
    run.font.size = Pt(14)
    run.font.color.rgb = GREY

    for _ in range(4):
        doc.add_paragraph("")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Prepared by Hunter Agent | InfoVision")
    run.font.size = Pt(10)
    run.font.color.rgb = GREY


def _add_executive_summary(
    doc: Document,
    client: str,
    competitors: list[str],
    category: str,
    geography: str,
    grouped: dict[str, list[dict]],
    all_results: list[dict],
) -> None:
    """Add executive summary with synthesized overview."""
    doc.add_heading("Executive Summary", level=1)

    total_sources = len(all_results)
    news_count = sum(1 for r in all_results if r.get("_type") == "news")
    web_count = total_sources - news_count
    credible_count = sum(1 for r in all_results if _is_credible(
        r.get("href") or r.get("url") or r.get("link", "")))
    comp_str = ", ".join(competitors) if competitors else "key competitors"
    sections_covered = [s for s in SECTION_ORDER if s in grouped]

    p = doc.add_paragraph()
    run = p.add_run(
        f"This report presents a comprehensive secondary research analysis of {client} "
        f"within the {category} sector ({geography} market focus). Research was conducted "
        f"on {date.today().strftime('%B %d, %Y')}, drawing from {total_sources} unique sources "
        f"({news_count} news articles, {web_count} web sources), of which {credible_count} "
        f"originate from tier-one credible publications. The analysis covers {client}'s "
        f"competitive positioning relative to {comp_str}, along with industry dynamics, "
        f"regulatory developments, and consumer behavior trends."
    )
    run.font.size = Pt(11)

    doc.add_heading("Coverage Overview", level=2)
    for section_title in sections_covered:
        results = grouped[section_title]
        cred_in_section = sum(1 for r in results if _is_credible(
            r.get("href") or r.get("url") or r.get("link", "")))
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(f"{section_title}: ")
        run.font.bold = True
        run.font.color.rgb = VIOLET
        p.add_run(f"{len(results)} sources ({cred_in_section} from credible outlets)")

    doc.add_heading("Research Scope", level=2)

    scope_items = [
        ("Primary Subject", client),
        ("Category", category),
        ("Geography", geography),
        ("Research Date", date.today().strftime("%B %d, %Y")),
    ]
    if competitors:
        scope_items.insert(1, ("Competitive Set", ", ".join(competitors)))

    for label, value in scope_items:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)


def _add_findings_section(
    doc: Document,
    title: str,
    results: list[dict],
    client: str,
) -> None:
    """Add a thematic findings section with source details and relevance callouts."""
    doc.add_heading(title, level=1)

    for i, r in enumerate(results):
        source_title = r.get("title") or "Untitled Source"
        url = r.get("href") or r.get("url") or r.get("link") or ""
        snippet = r.get("body") or r.get("snippet") or r.get("description") or ""
        source_date = r.get("date") or ""
        source_name = r.get("source") or _extract_domain(url)
        credible = _is_credible(url)

        heading_prefix = f"{i + 1}. "
        if credible:
            heading_prefix += "[Credible Source] "
        doc.add_heading(f"{heading_prefix}{source_title}", level=2)

        if source_name or source_date:
            p = doc.add_paragraph()
            meta_parts = []
            if source_name:
                meta_parts.append(source_name)
            if source_date:
                meta_parts.append(source_date)
            run = p.add_run(" | ".join(meta_parts))
            run.font.size = Pt(9)
            run.font.color.rgb = GREY
            run.font.italic = True

        if snippet:
            p = doc.add_paragraph()
            p.add_run(snippet)

        if url:
            p = doc.add_paragraph()
            run = p.add_run(f"Source: {url}")
            run.font.size = Pt(9)
            run.font.color.rgb = BLUE
            run.font.underline = True

        relevance = _generate_relevance(title, source_title, snippet, client)
        if relevance:
            p = doc.add_paragraph()
            run = p.add_run("Why This Matters: ")
            run.font.bold = True
            run.font.color.rgb = VIOLET
            run.font.size = Pt(10)
            run2 = p.add_run(relevance)
            run2.font.size = Pt(10)
            run2.font.italic = True


def _generate_relevance(section: str, title: str, snippet: str, client: str) -> str:
    """Generate a brief relevance statement based on section context."""
    section_relevance = {
        SECTION_NEWS: f"Recent developments that may impact {client}'s market position and stakeholder perception.",
        SECTION_INDUSTRY: f"Macro industry shifts that shape the competitive environment {client} operates in.",
        SECTION_COMPETITIVE: f"Competitive moves that may require a strategic response from {client}.",
        SECTION_CONSUMER: f"Evolving consumer expectations that could influence {client}'s product and messaging strategy.",
        SECTION_REGULATORY: f"Regulatory changes that may affect {client}'s operations, compliance, or market access.",
        SECTION_MARKET_DATA: f"Market sizing and projections relevant to {client}'s growth planning and investor narrative.",
    }
    return section_relevance.get(section, "")


def _add_opportunities_risks(
    doc: Document,
    client: str,
    category: str,
    grouped: dict[str, list[dict]],
) -> None:
    """Add opportunities and risks section based on research findings."""
    doc.add_heading("Opportunities & Risks", level=1)

    p = doc.add_paragraph()
    p.add_run(
        f"Based on the secondary research findings, the following strategic considerations "
        f"emerge for {client} in the {category} space."
    )

    doc.add_heading("Potential Opportunities", level=2)
    opportunities = []
    if SECTION_INDUSTRY in grouped:
        opportunities.append(
            f"Industry trends in {category} suggest evolving market dynamics that "
            f"{client} may capitalize on through early positioning."
        )
    if SECTION_CONSUMER in grouped:
        opportunities.append(
            f"Shifting consumer behaviors identified in the research may open new "
            f"engagement channels or product opportunities for {client}."
        )
    if SECTION_MARKET_DATA in grouped:
        opportunities.append(
            f"Market growth projections indicate expanding addressable market that "
            f"could support {client}'s revenue growth objectives."
        )
    if not opportunities:
        opportunities.append("Further primary research recommended to identify specific opportunities.")

    for opp in opportunities:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(opp)

    doc.add_heading("Potential Risks", level=2)
    risks = []
    if SECTION_COMPETITIVE in grouped:
        risks.append(
            f"Competitive activity suggests rivals are actively investing in the space, "
            f"which could pressure {client}'s market share."
        )
    if SECTION_REGULATORY in grouped:
        risks.append(
            f"Regulatory developments may introduce new compliance requirements or "
            f"constraints that affect {client}'s operations."
        )
    if SECTION_NEWS in grouped:
        risks.append(
            f"Recent news coverage warrants monitoring for potential reputational "
            f"or operational implications."
        )
    if not risks:
        risks.append("No immediate risk signals detected; continued monitoring recommended.")

    for risk in risks:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(risk)

    p = doc.add_paragraph()
    run = p.add_run(
        "Note: These opportunities and risks are directional indicators derived from "
        "secondary sources. Validation through primary research and internal data analysis "
        "is recommended before strategic action."
    )
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = GREY


def _add_source_appendix(
    doc: Document,
    all_results: list[dict],
    queries_run: list[str],
) -> None:
    """Add source citation appendix and methodology notes."""
    doc.add_heading("Appendix: Source Citations & Methodology", level=1)

    doc.add_heading("All Sources Referenced", level=2)
    sorted_results = _sort_by_credibility(all_results)

    for i, r in enumerate(sorted_results):
        title = r.get("title") or "Untitled"
        url = r.get("href") or r.get("url") or r.get("link") or ""
        source = r.get("source") or _extract_domain(url)
        source_date = r.get("date") or ""
        credible = _is_credible(url)

        p = doc.add_paragraph(style="List Number")
        run = p.add_run(title)
        run.font.bold = True
        run.font.size = Pt(10)

        if credible:
            badge = p.add_run(" [Credible]")
            badge.font.size = Pt(8)
            badge.font.color.rgb = VIOLET
            badge.font.bold = True

        meta_line = []
        if source:
            meta_line.append(source)
        if source_date:
            meta_line.append(source_date)
        if meta_line:
            p2 = doc.add_paragraph()
            run2 = p2.add_run("   " + " | ".join(meta_line))
            run2.font.size = Pt(9)
            run2.font.color.rgb = GREY

        if url:
            p3 = doc.add_paragraph()
            run3 = p3.add_run("   " + url)
            run3.font.size = Pt(9)
            run3.font.color.rgb = BLUE
            run3.font.underline = True

    doc.add_heading("Search Methodology", level=2)
    p = doc.add_paragraph()
    p.add_run(
        f"Research conducted on {date.today().strftime('%B %d, %Y')} using automated web and news "
        f"search via DuckDuckGo. A total of {len(queries_run)} context-aware search queries were "
        f"constructed to cover company news, industry trends, competitive landscape, consumer "
        f"behavior, regulatory developments, and market data. Results were deduplicated and "
        f"prioritized by source credibility ({len(CREDIBLE_DOMAINS)} credible domains tracked)."
    )

    doc.add_heading("Search Queries Used", level=3)
    for q in queries_run:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(q)
        run.font.size = Pt(10)
        run.font.italic = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("\n\n--- End of Report ---")
    run.font.size = Pt(10)
    run.font.color.rgb = GREY
    run.font.italic = True
