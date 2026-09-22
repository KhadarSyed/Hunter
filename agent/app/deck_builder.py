"""Assembles the final client deck from a DeckPlan: the branded template
skeleton (cover/TOC/objective/summary/thank-you + one divider per section),
historical slides copied in whole via COM for fidelity, and freshly drafted
slides built with python-pptx for content gaps the repository didn't cover.

Two-phase design to avoid ever mixing COM and python-pptx index math:
  Phase 1 (COM):        all structural changes - delete template's example
                         content slides, duplicate dividers, paste in reused
                         historical slides - in their final left-to-right order.
  Phase 2 (python-pptx): reopen the saved file and only ever append + reorder
                         (never delete), then edit text in place.
"""
from __future__ import annotations

import shutil
from datetime import date
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .slide_copy_com import DeckSurgeon

VIOLET = RGBColor(0x5B, 0x2C, 0x9D)
GREY = RGBColor(0x6E, 0x6E, 0x6E)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)

# Fixed layout of the source hunter_template.pptx (see references/slide_catalog.md
# and paragraph_index_map.md in the hunter-report-builder skill).
TEMPLATE_CONTENT_LIBRARY_INDICES = [6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16]
TEMPLATE_DIVIDER_INDEX = 5
TEMPLATE_SLIDE_COUNT = 17


def _flatten_paragraphs(slide):
    """All paragraphs across all text-bearing shapes on a slide, in shape order
    - matches how the skill's own paragraph-index map was derived."""
    paras = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            paras.extend(shape.text_frame.paragraphs)
    return paras


def _set_paragraph_text(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run()
    paragraph.runs[0].text = text
    for extra in paragraph.runs[1:]:
        extra.text = ""


def _move_slide(prs, old_index: int, new_index: int) -> None:
    """0-based indices. Reorders the slide XML in place."""
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    xml_slides.remove(slides[old_index])
    xml_slides.insert(new_index, slides[old_index])


def _blank_layout(prs):
    for layout in prs.slide_masters[0].slide_layouts:
        if layout.name.strip().lower() == "blank":
            return layout
    return prs.slide_masters[0].slide_layouts[-1]


def _add_drafted_slide(prs, eyebrow: str, title: str, paragraphs: list[str], source_note: str):
    """Builds a slide in the house style observed in real past decks: plain
    white background, small violet eyebrow label, bold title, prose body,
    grey footer - see brief_parser/deck_index sample inspection."""
    slide = prs.slides.add_slide(_blank_layout(prs))
    width = prs.slide_width
    height = prs.slide_height

    eyebrow_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), width - Inches(1.2), Inches(0.4))
    tf = eyebrow_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = eyebrow.upper()
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = VIOLET

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.75), width - Inches(1.2), Inches(0.9))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = BLACK

    body_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.7), width - Inches(1.2), height - Inches(2.6))
    tf = body_box.text_frame
    tf.word_wrap = True
    flat_paragraphs = []
    for item in paragraphs:
        if isinstance(item, list):
            flat_paragraphs.extend(str(x) for x in item)
        elif isinstance(item, dict):
            flat_paragraphs.append(str(item.get("text", item)))
        else:
            flat_paragraphs.append(str(item))
    for i, para_text in enumerate(flat_paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = para_text
        run.font.size = Pt(14)
        run.font.color.rgb = BLACK
        p.space_after = Pt(10)

    footer_box = slide.shapes.add_textbox(Inches(0.6), height - Inches(0.55), width - Inches(1.2), Inches(0.35))
    tf = footer_box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = source_note
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = GREY
    p.alignment = PP_ALIGN.RIGHT

    return slide


def _add_divider_slide(prs, title: str, parenthetical: str = ""):
    """Creates a section divider slide with violet styling."""
    slide = prs.slides.add_slide(_blank_layout(prs))
    width = prs.slide_width
    height = prs.slide_height

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.5), width - Inches(1.6), Inches(1.2))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.color.rgb = VIOLET
    p.alignment = PP_ALIGN.LEFT

    if parenthetical:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.7), width - Inches(1.6), Inches(0.5))
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = parenthetical
        run.font.size = Pt(16)
        run.font.italic = True
        run.font.color.rgb = GREY

    return slide


def _tag_reused_slide(slide, width, reference_note: str) -> None:
    """Adds a small reference label to a historical slide copied in whole,
    so it's clear in the finished deck which past project it came from -
    the underlying chart/image content is left completely untouched."""
    box = slide.shapes.add_textbox(Inches(0.3), Inches(0.05), width - Inches(0.6), Inches(0.3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = reference_note
    run.font.size = Pt(8)
    run.font.italic = True
    run.font.color.rgb = GREY


def build_skeleton(output_path: str, sections: list[dict]) -> dict:
    """Phase 1 (COM): strip the template's example content library, create one
    divider per section, paste in every reused historical slide, in final
    order. Returns a plan of recorded absolute slide positions for phase 2.
    """
    layout: dict[str, Any] = {"dividers": [], "sections": []}

    with DeckSurgeon() as surgeon:
        surgeon.open_dest(output_path)

        for idx in sorted(TEMPLATE_CONTENT_LIBRARY_INDICES, reverse=True):
            surgeon.delete_slide(idx)

        n = len(sections)
        for _ in range(n - 1):
            surgeon.duplicate_slide(TEMPLATE_DIVIDER_INDEX)

        shift = 0
        for i, section in enumerate(sections):
            divider_idx = TEMPLATE_DIVIDER_INDEX + i + shift
            layout["dividers"].append(divider_idx)
            insert_at = divider_idx + 1
            reused_positions = []
            for j, slide_ref in enumerate(section.get("reused", [])):
                pos = surgeon.paste_slide(slide_ref["deck_path"], slide_ref["slide_no"], insert_at + j)
                reused_positions.append(pos)
            shift += len(section.get("reused", []))
            layout["sections"].append({
                "divider_index": divider_idx,
                "reused_positions": reused_positions,
            })

        surgeon.save()

    return layout


def apply_content(output_path: str, plan: dict, layout: dict) -> None:
    """Phase 2 (python-pptx): fill in cover/TOC/objective/summary/dividers,
    tag reused slides with their provenance, append + reorder drafted slides.
    """
    prs = Presentation(output_path)

    _fill_cover(prs, plan["cover"])
    _fill_toc(prs, [s["title"] for s in plan["sections"]])
    _fill_objective(prs, plan["objective"])
    _fill_summary(prs, plan["summary"])

    for i, section in enumerate(plan["sections"]):
        divider_idx = layout["sections"][i]["divider_index"]
        slide = prs.slides[divider_idx - 1]
        paras = _flatten_paragraphs(slide)
        if len(paras) >= 2:
            _set_paragraph_text(paras[0], section["title"])
            _set_paragraph_text(paras[1], section.get("parenthetical", ""))

        for pos, slide_ref in zip(layout["sections"][i]["reused_positions"], section.get("reused", [])):
            reused_slide = prs.slides[pos - 1]
            note = f"REFERENCE: {slide_ref.get('source_client', 'past project')} · {slide_ref.get('source_date', '')}"
            _tag_reused_slide(reused_slide, prs.slide_width, note.strip(" ·"))

    phase2_shift = 0
    today = date.today().strftime("%B %Y")
    for i, section in enumerate(plan["sections"]):
        drafted_list = section.get("drafted", [])
        if not drafted_list:
            continue
        base_position = (
            layout["sections"][i]["divider_index"]
            + len(section.get("reused", []))
            + phase2_shift
            + 1
        )
        for k, drafted in enumerate(drafted_list):
            slide = _add_drafted_slide(
                prs,
                eyebrow=drafted.get("eyebrow", section["title"]),
                title=drafted["title"],
                paragraphs=drafted["paragraphs"],
                source_note=f"SOURCE: AGENT DRAFT | {today}",
            )
            old_index = len(prs.slides) - 1
            new_index = (base_position + k) - 1  # 0-based target
            _move_slide(prs, old_index, new_index)
            phase2_shift += 1

    prs.save(output_path)


def _ensure_skeleton() -> Path:
    """Return the path to the pre-built skeleton (template with content library
    slides 6-16 already deleted). Builds it on first use via python-pptx."""
    skeleton = Path(__file__).resolve().parent.parent / "data" / "hunter_template_skeleton.pptx"
    if skeleton.exists():
        return skeleton
    from pptx.oxml.ns import qn as _qn
    src = Path(__file__).resolve().parent.parent / "data" / "hunter_template.pptx"
    prs = Presentation(str(src))
    sldIdLst = prs.slides._sldIdLst
    for idx in range(15, 4, -1):
        elements = list(sldIdLst)
        rId = elements[idx].get(_qn("r:id"))
        sldIdLst.remove(elements[idx])
        prs.part.drop_rel(rId)
    prs.save(str(skeleton))
    return skeleton


def build_template_deck(template_path: str, output_path: str, plan: dict) -> None:
    """Builds a template deck with visual slides (charts, tables, matrices).
    Uses the pre-built skeleton + chart_builders for rich visual output."""
    import tempfile
    import uuid
    from .chart_builders import build_visual_slide

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    tmp_dir = Path(tempfile.gettempdir()) / "hunter_agent_build"
    tmp_dir.mkdir(exist_ok=True)
    unique = uuid.uuid4().hex[:8]
    tmp_path = str(tmp_dir / f"build_{unique}.pptx")
    shutil.copy(str(_ensure_skeleton()), tmp_path)

    prs = Presentation(tmp_path)

    _fill_cover(prs, plan["cover"])
    section_titles = [s["title"] for s in plan["sections"]]
    _fill_toc(prs, section_titles)
    _fill_objective(prs, plan["objective"])
    _fill_summary_with_data(prs, plan.get("summary", {}))

    divider_paras = _flatten_paragraphs(prs.slides[4])
    today = date.today().strftime("%B %Y")

    for i, section in enumerate(plan["sections"]):
        if i == 0:
            if len(divider_paras) >= 2:
                _set_paragraph_text(divider_paras[0], section["title"])
                _set_paragraph_text(divider_paras[1], section.get("parenthetical", ""))
        else:
            _add_divider_slide(prs, section["title"], section.get("parenthetical", ""))

        for drafted in section.get("drafted", []):
            if drafted.get("visual_type"):
                build_visual_slide(prs, drafted)
            else:
                _add_drafted_slide(
                    prs,
                    eyebrow=drafted.get("eyebrow", section["title"]),
                    title=drafted["title"],
                    paragraphs=drafted.get("paragraphs", []),
                    source_note=f"SOURCE: TEMPLATE DRAFT | {today}",
                )

    _move_slide(prs, 5, len(prs.slides) - 1)

    prs.save(tmp_path)
    shutil.copy2(tmp_path, output_path)
    try:
        Path(tmp_path).unlink()
    except OSError:
        pass


def build_deck(template_path: str, output_path: str, plan: dict) -> None:
    import tempfile
    import uuid
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    tmp_dir = Path(tempfile.gettempdir()) / "hunter_agent_build"
    tmp_dir.mkdir(exist_ok=True)
    unique = uuid.uuid4().hex[:8]
    tmp_path = str(tmp_dir / f"build_{unique}.pptx")
    shutil.copy(template_path, tmp_path)
    layout = build_skeleton(tmp_path, plan["sections"])
    apply_content(tmp_path, plan, layout)
    shutil.copy2(tmp_path, output_path)
    try:
        Path(tmp_path).unlink()
    except OSError:
        pass


# ---- fixed-frame field fillers (indices from the skill's paragraph_index_map.md) ----

def _fill_cover(prs, cover: dict) -> None:
    slide = prs.slides[0]
    paras = _flatten_paragraphs(slide)
    if len(paras) > 3:
        _set_paragraph_text(paras[1], cover.get("report_type", "Research Report"))
        _set_paragraph_text(paras[2], cover.get("subject", ""))
        _set_paragraph_text(paras[3], cover.get("date", ""))


def _fill_toc(prs, section_titles: list[str]) -> None:
    slide = prs.slides[1]
    paras = _flatten_paragraphs(slide)
    slots = paras[2:15] if len(paras) > 15 else paras[2:]
    for i, slot in enumerate(slots):
        _set_paragraph_text(slot, section_titles[i] if i < len(section_titles) else "")


def _fill_objective(prs, objective: dict) -> None:
    slide = prs.slides[2]
    paras = _flatten_paragraphs(slide)
    mapping = {
        2: objective.get("geography", "[INSERT GEOGRAPHY]"),
        3: objective.get("time_period", "[INSERT TIME PERIOD]"),
        4: objective.get("sources_tools", "[INSERT SOURCES / TOOLS]"),
        6: objective.get("research_objective", ""),
        14: objective.get("competitive_research", "[INSERT COMPETITIVE SET]"),
        16: objective.get("target_audience", "[INSERT TARGET AUDIENCE]"),
    }
    for i, focus_bullet in enumerate(objective.get("focus_bullets", [])[:3]):
        mapping[10 + i] = focus_bullet
    for idx, text in mapping.items():
        if idx < len(paras):
            if isinstance(text, list):
                text = ", ".join(str(t) for t in text)
            _set_paragraph_text(paras[idx], str(text))


def _fill_summary_with_data(prs, summary: dict) -> None:
    """Fills the summary slide with realistic dummy data from template mode."""
    slide = prs.slides[3]
    paras = _flatten_paragraphs(slide)
    mapping = {
        2: summary.get("share_of_voice", "[INSERT SHARE OF VOICE]"),
        3: summary.get("sentiment_split", "[INSERT SENTIMENT SPLIT]"),
        4: summary.get("key_themes", "[INSERT KEY THEMES]"),
        5: summary.get("top_hashtags", "[INSERT TOP HASHTAGS]"),
        6: summary.get("seasonal_peaks", "[INSERT SEASONAL PEAKS / MOMENTS]"),
        9: summary.get("client", ""),
        10: summary.get("category", ""),
        11: summary.get("competitors", "[INSERT COMPETITORS]"),
    }
    profile_bullets = summary.get("consumer_profile_bullets", [])
    for i in range(13, 18):
        mapping[i] = profile_bullets[i - 13] if i - 13 < len(profile_bullets) else "[INSERT CONSUMER PROFILE INSIGHT]"
    for idx, text in mapping.items():
        if idx < len(paras):
            if isinstance(text, list):
                text = ", ".join(str(t) for t in text)
            _set_paragraph_text(paras[idx], str(text))


def _fill_summary(prs, summary: dict) -> None:
    slide = prs.slides[3]
    paras = _flatten_paragraphs(slide)
    mapping = {
        2: "[INSERT SHARE OF VOICE]",
        3: "[INSERT SENTIMENT SPLIT]",
        4: "[INSERT KEY THEMES]",
        5: "[INSERT TOP HASHTAGS]",
        6: "[INSERT SEASONAL PEAKS / MOMENTS]",
        9: summary.get("client", ""),
        10: summary.get("category", ""),
        11: summary.get("competitors", "[INSERT COMPETITORS]"),
    }
    profile_bullets = summary.get("consumer_profile_bullets", [])
    for i in range(13, 18):
        mapping[i] = profile_bullets[i - 13] if i - 13 < len(profile_bullets) else "[INSERT CONSUMER PROFILE INSIGHT]"
    for idx, text in mapping.items():
        if idx < len(paras):
            _set_paragraph_text(paras[idx], text)
