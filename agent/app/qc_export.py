"""Export QC results as annotated Excel files."""
from __future__ import annotations

import logging
import time
from pathlib import Path

from . import intelligence_store as store
from . import qc_parser

logger = logging.getLogger(__name__)

EXPORT_DIR = Path(__file__).resolve().parent.parent / "data" / "exports"


def export_qc_results(run_id: int) -> dict:
    """Generate an annotated Excel export with QC findings."""
    run = store.get_qc_run(run_id)
    if not run:
        return {"error": "QC run not found"}

    report = store.get_qc_report(run["report_id"])
    if not report:
        return {"error": "Report not found"}

    findings = store.get_qc_findings(run_id)
    mapping = report.get("field_mapping") or {}

    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    except ImportError:
        return {"error": "openpyxl not installed"}

    rows = qc_parser.get_rows_from_file(report["file_path"], mapping)

    row_findings: dict[int, list[dict]] = {}
    for f in findings:
        rn = f.get("row_number", 0)
        row_findings.setdefault(rn, []).append(f)

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    base_name = Path(report["file_name"]).stem
    export_name = f"{base_name}_QC_Report.xlsx"
    export_path = EXPORT_DIR / export_name

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "QC Results"

    header_fill = PatternFill(start_color="0F7B6C", end_color="0F7B6C", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    severity_fills = {
        "critical": PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid"),
        "high": PatternFill(start_color="FFE0CC", end_color="FFE0CC", fill_type="solid"),
        "medium": PatternFill(start_color="FFFACC", end_color="FFFACC", fill_type="solid"),
        "low": PatternFill(start_color="E6F5E6", end_color="E6F5E6", fill_type="solid"),
        "info": PatternFill(start_color="E6F0FF", end_color="E6F0FF", fill_type="solid"),
    }

    original_cols = list(report.get("columns_json") or [])
    qc_cols = ["QC_Score", "QC_Severity", "QC_Issues", "QC_Details"]
    all_cols = original_cols + qc_cols

    for col_idx, col_name in enumerate(all_cols, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    original_data = qc_parser.parse_file(report["file_path"])
    raw_preview = original_data.get("preview", []) if "error" not in original_data else []

    all_raw_rows = _read_all_original(report["file_path"], original_cols)

    for row_idx, raw_row in enumerate(all_raw_rows, 2):
        for col_idx, col_name in enumerate(original_cols, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=raw_row.get(col_name, ""))
            cell.border = thin_border

        row_num = row_idx
        issues = row_findings.get(row_num, [])

        if issues:
            max_severity = _max_severity(issues)
            score = max(0, 100 - sum(
                {"critical": 25, "high": 15, "medium": 8, "low": 3, "info": 0}.get(f["severity"], 0)
                for f in issues
            ))
            issue_types = ", ".join(sorted(set(f["check_type"] for f in issues)))
            details = " | ".join(f["message"] for f in issues[:5])

            score_cell = ws.cell(row=row_idx, column=len(original_cols) + 1, value=score)
            sev_cell = ws.cell(row=row_idx, column=len(original_cols) + 2, value=max_severity.upper())
            issues_cell = ws.cell(row=row_idx, column=len(original_cols) + 3, value=issue_types)
            details_cell = ws.cell(row=row_idx, column=len(original_cols) + 4, value=details)

            fill = severity_fills.get(max_severity)
            if fill:
                for c in (score_cell, sev_cell, issues_cell, details_cell):
                    c.fill = fill
                    c.border = thin_border
        else:
            ws.cell(row=row_idx, column=len(original_cols) + 1, value=100).border = thin_border
            ws.cell(row=row_idx, column=len(original_cols) + 2, value="PASS").border = thin_border

    for col_idx in range(1, len(all_cols) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 18

    ws_summary = wb.create_sheet("QC Summary")
    summary_data = [
        ("Report", report["file_name"]),
        ("Total Rows", report["row_count"]),
        ("Total Findings", len(findings)),
        ("Overall Score", run.get("score", "N/A")),
    ]

    severity_counts = {}
    for f in findings:
        sev = f.get("severity", "info")
        severity_counts[sev] = severity_counts.get(sev, 0) + 1

    summary_data.append(("", ""))
    summary_data.append(("Severity Breakdown", ""))
    for sev in ("critical", "high", "medium", "low", "info"):
        if severity_counts.get(sev, 0):
            summary_data.append((sev.upper(), severity_counts[sev]))

    for row_idx, (label, value) in enumerate(summary_data, 1):
        ws_summary.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
        ws_summary.cell(row=row_idx, column=2, value=value)

    ws_summary.column_dimensions["A"].width = 25
    ws_summary.column_dimensions["B"].width = 30

    wb.save(str(export_path))

    export_id = _save_export_record(run_id, export_name, str(export_path))

    return {
        "export_id": export_id,
        "file_name": export_name,
        "file_path": str(export_path),
    }


def _max_severity(issues: list[dict]) -> str:
    order = ["critical", "high", "medium", "low", "info"]
    for sev in order:
        if any(f.get("severity") == sev for f in issues):
            return sev
    return "info"


def _read_all_original(file_path: str, columns: list[str]) -> list[dict]:
    """Read original file preserving original column names."""
    from pathlib import Path
    ext = Path(file_path).suffix.lower()

    if ext == ".csv":
        return qc_parser._read_all_csv(Path(file_path))
    elif ext in (".xlsx", ".xls"):
        return qc_parser._read_all_excel(Path(file_path))
    elif ext == ".pptx":
        return qc_parser._read_all_pptx(Path(file_path))
    elif ext == ".docx":
        return qc_parser._read_all_docx(Path(file_path))
    return []


def _save_export_record(run_id: int, file_name: str, file_path: str) -> int:
    conn = store._conn()
    now = time.time()
    cur = conn.execute(
        "INSERT INTO qc_exports (run_id, file_name, file_path, format, created_at) VALUES (?, ?, ?, 'xlsx', ?)",
        (run_id, file_name, file_path, now),
    )
    eid = cur.lastrowid
    conn.commit()
    conn.close()
    return eid
