"""Build and incrementally refresh the searchable slide index over the repository
of past decks: per-slide title/body text, chart/picture presence, and an
embedding vector, persisted in memory.slide_index.
"""
from __future__ import annotations

import fnmatch
import re
from pathlib import Path
from typing import Callable, Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from . import memory
from .config import Settings
from .ollama_client import OllamaClient

EventFn = Optional[Callable[[str, dict], None]]

_CLIENT_PREFIXES = ["Hunter PR Research_", "Hunter PR Research "]


def guess_client_from_filename(filename: str) -> str:
    stem = Path(filename).stem
    for prefix in _CLIENT_PREFIXES:
        if stem.startswith(prefix):
            stem = stem[len(prefix):]
            break
    # Client name is the token up to the next separator (_ or -)
    m = re.split(r"[_\-]", stem, maxsplit=1)
    token = m[0].strip() if m else stem.strip()
    return token or "Unknown"


_MONTH_YEAR_RE = re.compile(
    r"(?:\d{1,2}\s+)?(January|February|March|April|May|June|July|August|September|October|November|December)"
    r"[a-z]*\.?,?\s+\d{4}",
    re.IGNORECASE,
)


def guess_date_from_filename(filename: str) -> str:
    stem = Path(filename).stem
    match = _MONTH_YEAR_RE.search(stem)
    return match.group(0) if match else ""


def file_fingerprint(path: Path) -> tuple[float, str]:
    stat = path.stat()
    return stat.st_mtime, f"{stat.st_size}:{int(stat.st_mtime)}"


def is_ignored(filename: str, ignore_patterns: list[str]) -> bool:
    return any(fnmatch.fnmatch(filename, pat) for pat in ignore_patterns)


def _slide_text(slide) -> tuple[str, str, bool, bool]:
    title = ""
    if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
        title = slide.shapes.title.text_frame.text.strip()

    body_lines = []
    has_chart = False
    has_picture = False
    for shape in slide.shapes:
        if getattr(shape, "has_chart", False):
            has_chart = True
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            has_picture = True
        if shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text and text != title:
                body_lines.append(text)
    if not title and body_lines:
        title = body_lines[0][:120]
    return title, "\n".join(body_lines), has_chart, has_picture


def index_deck(deck_path: Path, ollama: OllamaClient) -> int:
    """(Re)index a single deck file. Returns number of slides indexed."""
    memory.delete_slides_for_deck(str(deck_path))
    mtime, fingerprint = file_fingerprint(deck_path)
    client_guess = guess_client_from_filename(deck_path.name)

    prs = Presentation(str(deck_path))
    count = 0
    for i, slide in enumerate(prs.slides, start=1):
        title, body, has_chart, has_picture = _slide_text(slide)
        combined = f"{title}\n{body}".strip()
        if not combined:
            continue
        vec = ollama.embed(combined)
        memory.upsert_slide(
            deck_path=str(deck_path),
            slide_no=i,
            title=title,
            body_text=body,
            has_chart=has_chart,
            has_picture=has_picture,
            client_guess=client_guess,
            embedding=vec.tobytes(),
            file_mtime=mtime,
            file_hash=fingerprint,
        )
        count += 1
    return count


def index_repository(settings: Settings, ollama: OllamaClient, on_event: EventFn = None) -> dict:
    repo = Path(settings.repo_dir)
    if not repo.exists():
        raise FileNotFoundError(f"Repository folder not found: {repo}")

    existing = memory.get_indexed_deck_hashes()
    deck_files = sorted(repo.glob("*.pptx"))
    seen_paths = set()

    stats = {"scanned": 0, "reindexed": 0, "skipped_unchanged": 0, "ignored": 0, "slides": 0}

    for deck_path in deck_files:
        seen_paths.add(str(deck_path))
        if is_ignored(deck_path.name, settings.ignore_patterns):
            stats["ignored"] += 1
            continue
        stats["scanned"] += 1
        mtime, fingerprint = file_fingerprint(deck_path)
        prev = existing.get(str(deck_path))
        if prev and prev[1] == fingerprint:
            stats["skipped_unchanged"] += 1
            continue

        if on_event:
            on_event("indexing_deck", {"deck": deck_path.name})
        try:
            n = index_deck(deck_path, ollama)
            stats["reindexed"] += 1
            stats["slides"] += n
            if on_event:
                on_event("indexed_deck", {"deck": deck_path.name, "slides": n})
        except Exception as e:  # keep indexing the rest even if one deck fails
            if on_event:
                on_event("index_error", {"deck": deck_path.name, "error": str(e)})

    # Drop slides for decks that no longer exist in the repo folder
    for deck_path_str in list(existing.keys()):
        if deck_path_str not in seen_paths:
            memory.delete_slides_for_deck(deck_path_str)

    stats["total_slides"] = memory.slide_count()
    stats["total_decks"] = memory.deck_count()
    return stats


if __name__ == "__main__":
    import sys
    from .config import load_settings

    settings = load_settings()
    ollama = OllamaClient(settings.ollama_host, settings.embed_model, settings.chat_model)
    if not ollama.is_reachable():
        print("Ollama is not reachable at", settings.ollama_host)
        sys.exit(1)
    memory.init_db()

    def log(event, payload):
        print(f"[{event}] {payload}")

    result = index_repository(settings, ollama, on_event=log)
    print("Done:", result)
