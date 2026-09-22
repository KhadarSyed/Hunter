"""Extract plain text from a client brief file (.docx, .txt, .pptx)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from pptx import Presentation

SUPPORTED_SUFFIXES = {".docx", ".txt", ".pptx"}


def parse_brief(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix == ".pptx":
        return _parse_pptx(path)
    raise ValueError(f"Unsupported brief format: {suffix}")


def _parse_docx(path: Path) -> str:
    doc = Document(str(path))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _parse_pptx(path: Path) -> str:
    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_lines = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                slide_lines.append(shape.text_frame.text.strip())
        if slide_lines:
            parts.append(f"[Slide {i}]\n" + "\n".join(slide_lines))
    return "\n\n".join(parts)


def guess_client_name(brief_text: str, filename: str) -> str:
    """Best-effort client name guess; refined later by the LLM reasoning pass."""
    stem = Path(filename).stem
    for sep in ("_", "-"):
        if sep in stem:
            token = stem.split(sep)[0].strip()
            if token and token.lower() not in ("brief", "client", "new"):
                return token
    return stem.strip() or "Unknown Client"
