"""Visual slide builders — consulting-quality, visual-first layouts.

Every builder creates a slide that leads with the visualization and uses
only short callout text (labels, KPI numbers, 1-sentence bullets).
No paragraphs.  Charts and infographics dominate the slide real estate.

Palette from Hunter PR reference decks:
  VIOLET  #5E35B1    RED     #DE2A00    BLUE    #A6CAEC
  LTPURP  #D1C4E9    GREEN   #4EA72E    TEAL    #156082
  PINK    #A02B93    DKGREEN #196B24
"""
from __future__ import annotations

from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

VIOLET = RGBColor(0x5E, 0x35, 0xB1)
RED = RGBColor(0xDE, 0x2A, 0x00)
BLUE = RGBColor(0xA6, 0xCA, 0xEC)
LT_PURPLE = RGBColor(0xD1, 0xC4, 0xE9)
GREEN = RGBColor(0x4E, 0xA7, 0x2E)
TEAL = RGBColor(0x15, 0x60, 0x82)
PINK = RGBColor(0xA0, 0x2B, 0x93)
DK_GREEN = RGBColor(0x19, 0x6B, 0x24)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x6E, 0x6E, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BRAND_VIOLET = RGBColor(0x5B, 0x2C, 0x9D)
CARD_BG = RGBColor(0xF5, 0xF0, 0xFA)

SERIES_PALETTE = [VIOLET, TEAL, PINK, GREEN, DK_GREEN, RED, BLUE, LT_PURPLE]


def _blank_layout(prs: Presentation):
    for layout in prs.slide_masters[0].slide_layouts:
        if layout.name.strip().lower() == "blank":
            return layout
    return prs.slide_masters[0].slide_layouts[-1]


def _add_eyebrow(slide, text, w):
    box = slide.shapes.add_textbox(Inches(0.3), Inches(0.15), w - Inches(0.6), Inches(0.22))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = text.upper()
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = BRAND_VIOLET


def _add_title(slide, text, w):
    box = slide.shapes.add_textbox(Inches(0.3), Inches(0.35), Inches(7.0), Inches(0.4))
    box.text_frame.word_wrap = True
    r = box.text_frame.paragraphs[0].add_run()
    r.text = text
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = BLACK


def _add_takeaway(slide, text, w):
    """Short key-takeaway callout in the top-right — 1-2 sentences max."""
    shape = slide.shapes.add_shape(5, Inches(7.6), Inches(0.15), Inches(5.5), Inches(0.55))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = LT_PURPLE
    shape.line.width = Pt(0.5)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.06)
    tf.margin_bottom = Inches(0.06)
    r = tf.paragraphs[0].add_run()
    r.text = text
    r.font.size = Pt(9)
    r.font.color.rgb = BLACK


def _add_source(slide, source, w, h):
    box = slide.shapes.add_textbox(Inches(0.3), h - Inches(0.22), w - Inches(0.6), Inches(0.18))
    r = box.text_frame.paragraphs[0].add_run()
    r.text = source
    r.font.size = Pt(7)
    r.font.italic = True
    r.font.color.rgb = GREY


def _add_bullets(slide, bullets, left, top, width, height):
    """Short bullet list — each item is one short sentence."""
    if not bullets:
        return
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets if isinstance(bullets, list) else [bullets]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = f"•  {b}" if not str(b).startswith("•") else str(b)
        r.font.size = Pt(9)
        r.font.color.rgb = BLACK
        p.space_after = Pt(3)


def _header(slide, data, w, h):
    _add_eyebrow(slide, data.get("eyebrow", ""), w)
    _add_title(slide, data.get("title", ""), w)
    if data.get("key_takeaway"):
        _add_takeaway(slide, data["key_takeaway"], w)
    _add_source(slide, data.get("source", ""), w, h)


def _style_chart(chart, legend=False):
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(8)
    try:
        chart.chart_style = 2
    except Exception:
        pass


# ═══════════════════════════════════════════════════════════════════════
# 1. KPI TILES — row of big-number metric cards
# ═══════════════════════════════════════════════════════════════════════

def add_kpi_tiles_slide(prs: Presentation, data: dict) -> Any:
    """Row of 3-5 KPI metric cards — big number, label, trend indicator.

    data: tiles: [{value, label, subtitle, trend_up?}, ...]
    """
    slide = prs.slides.add_slide(_blank_layout(prs))
    w, h = prs.slide_width, prs.slide_height
    _header(slide, data, w, h)

    tiles = data.get("tiles", [])
    n = len(tiles)
    if n == 0:
        return slide

    tile_w = min(2.8, 12.0 / n)
    gap = 0.2
    total = n * tile_w + (n - 1) * gap
    start_x = (13.33 - total) / 2
    top = Inches(1.2)
    tile_h = Inches(2.8)

    for i, tile in enumerate(tiles):
        x = Inches(start_x + i * (tile_w + gap))
        shape = slide.shapes.add_shape(5, x, top, Inches(tile_w), tile_h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = LT_PURPLE
        shape.line.width = Pt(0.5)

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.15)
        tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(tile.get("value", "—"))
        r.font.size = Pt(32)
        r.font.bold = True
        r.font.color.rgb = SERIES_PALETTE[i % len(SERIES_PALETTE)]

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = tile.get("label", "")
        r2.font.size = Pt(10)
        r2.font.bold = True
        r2.font.color.rgb = BLACK

        if tile.get("subtitle"):
            p3 = tf.add_paragraph()
            p3.alignment = PP_ALIGN.CENTER
            p3.space_before = Pt(6)
            r3 = p3.add_run()
            trend = tile.get("trend_up")
            arrow = "▲ " if trend is True else ("▼ " if trend is False else "")
            r3.text = f"{arrow}{tile['subtitle']}"
            r3.font.size = Pt(8)
            r3.font.color.rgb = GREEN if trend is True else (RED if trend is False else GREY)

    if data.get("bullets"):
        _add_bullets(slide, data["bullets"], Inches(0.4), Inches(4.3), w - Inches(0.8), Inches(2.8))

    return slide


# ═══════════════════════════════════════════════════════════════════════
# 2. DOUGHNUT SENTIMENT GRID
# ═══════════════════════════════════════════════════════════════════════

def add_doughnut_sentiment_slide(prs: Presentation, data: dict) -> Any:
    slide = prs.slides.add_slide(_blank_layout(prs))
    w, h = prs.slide_width, prs.slide_height
    _header(slide, data, w, h)

    brands = data.get("brands", [])
    n = len(brands)
    if n == 0:
        return slide

    donut_w = Inches(2.4)
    donut_h = Inches(2.4)
    total_w = n * 3.2
    start_x = (13.33 - total_w) / 2
    top = Inches(1.2)

    for i, brand in enumerate(brands):
        x = Inches(start_x + i * 3.2)
        pos, neg, neu = brand.get("positive", 50), brand.get("negative", 20), brand.get("neutral", 30)

        cd = CategoryChartData()
        cd.categories = ["Positive", "Negative", "Neutral"]
        cd.add_series("Sentiment", (pos, neg, neu))

        cf = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, x, top, donut_w, donut_h, cd)
        chart = cf.chart
        _style_chart(chart)
        chart.plots[0].has_data_labels = False
        for pi, color in enumerate([VIOLET, RED, BLUE]):
            chart.series[0].points[pi].format.fill.solid()
            chart.series[0].points[pi].format.fill.fore_color.rgb = color

        nb = slide.shapes.add_textbox(x, top + donut_h + Inches(0.05), donut_w, Inches(0.25))
        nb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = nb.text_frame.paragraphs[0].add_run()
        r.text = brand["name"]
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = BLACK

        pb = slide.shapes.add_textbox(x, top + donut_h + Inches(0.3), donut_w, Inches(0.2))
        pb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = pb.text_frame.paragraphs[0].add_run()
        r.text = f"+{pos}%  -{neg}%  ~{neu}%"
        r.font.size = Pt(8)
        r.font.color.rgb = GREY

    legend_top = top + donut_h + Inches(0.65)
    for ci, (label, color) in enumerate([("Positive", VIOLET), ("Negative", RED), ("Neutral", BLUE)]):
        lx = Inches(start_x + ci * 2.2)
        dot = slide.shapes.add_shape(1, lx, legend_top, Inches(0.12), Inches(0.12))
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.fill.background()
        lb = slide.shapes.add_textbox(lx + Inches(0.18), legend_top - Inches(0.02), Inches(1.0), Inches(0.18))
        lb.text_frame.paragraphs[0].add_run().text = label
        lb.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
        lb.text_frame.paragraphs[0].runs[0].font.color.rgb = GREY

    if data.get("bullets"):
        _add_bullets(slide, data["bullets"], Inches(0.4), Inches(5.0), w - Inches(0.8), Inches(2.0))

    return slide


# ═══════════════════════════════════════════════════════════════════════
# 3. DOUGHNUT GAUGES — SOV / KPI indicators
# ═══════════════════════════════════════════════════════════════════════

def add_doughnut_gauge_slide(prs: Presentation, data: dict) -> Any:
    slide = prs.slides.add_slide(_blank_layout(prs))
    w, h = prs.slide_width, prs.slide_height
    _header(slide, data, w, h)

    gauges = data.get("gauges", [])
    n = len(gauges)
    if n == 0:
        return slide

    gauge_sz = Inches(2.8)
    total_w = n * 3.5
    start_x = (13.33 - total_w) / 2
    top = Inches(1.3)

    for i, g in enumerate(gauges):
        x = Inches(start_x + i * 3.5)
        val = g.get("value", 0)
        cd = CategoryChartData()
        cd.categories = ["Fill", "Empty"]
        cd.add_series("SOV", (val, 100 - val))
        cf = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, x, top, gauge_sz, gauge_sz, cd)
        chart = cf.chart
        _style_chart(chart)
        chart.plots[0].has_data_labels = False
        chart.series[0].points[0].format.fill.solid()
        chart.series[0].points[0].format.fill.fore_color.rgb = SERIES_PALETTE[i % len(SERIES_PALETTE)]
        chart.series[0].points[1].format.fill.solid()
        chart.series[0].points[1].format.fill.fore_color.rgb = LT_PURPLE

        pb = slide.shapes.add_textbox(x, top + Inches(0.9), gauge_sz, Inches(0.5))
        pb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = pb.text_frame.paragraphs[0].add_run()
        r.text = f"{val}%"
        r.font.size = Pt(24)
        r.font.bold = True
        r.font.color.rgb = SERIES_PALETTE[i % len(SERIES_PALETTE)]

        nb = slide.shapes.add_textbox(x, top + gauge_sz + Inches(0.05), gauge_sz, Inches(0.25))
        nb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = nb.text_frame.paragraphs[0].add_run()
        r.text = g["name"]
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = BLACK

        if g.get("label"):
            lb = slide.shapes.add_textbox(x, top + gauge_sz + Inches(0.3), gauge_sz, Inches(0.2))
            lb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            r = lb.text_frame.paragraphs[0].add_run()
            r.text = g["label"]
            r.font.size = Pt(8)
            r.font.color.rgb = GREY

    if data.get("bullets"):
        _add_bullets(slide, data["bullets"], Inches(0.4), Inches(5.0), w - Inches(0.8), Inches(2.0))

    return slide


# ═══════════════════════════════════════════════════════════════════════
# 4. LINE TREND
# ═══════════════════════════════════════════════════════════════════════

def add_line_trend_slide(prs: Presentation, data: dict) -> Any:
    slide = prs.slides.add_slide(_blank_layout(prs))
    w, h = prs.slide_width, prs.slide_height
    _header(slide, data, w, h)

    cd = CategoryChartData()
    cd.categories = data.get("categories", [])
    for s in data.get("series", []):
        cd.add_series(s["name"], s["values"])

    cf = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE, Inches(0.3), Inches(1.0), Inches(12.7), Inches(4.5), cd
    )
    chart = cf.chart
    _style_chart(chart, legend=True)

    for si, series in enumerate(chart.series):
        series.format.line.color.rgb = SERIES_PALETTE[si % len(SERIES_PALETTE)]
        series.format.line.width = Pt(2.5)
        series.smooth = True

    try:
        chart.category_axis.tick_labels.font.size = Pt(8)
        chart.value_axis.tick_labels.font.size = Pt(8)
        chart.value_axis.has_major_gridlines = True
    except Exception:
        pass

    if data.get("bullets"):
        _add_bullets(slide, data["bullets"], Inches(0.4), Inches(5.7), w - Inches(0.8), Inches(1.3))

    return slide


# ═══════════════════════════════════════════════════════════════════════
# 5. BAR CHART WITH CALLOUT CARDS
# ═══════════════════════════════════════════════════════════════════════

def add_bar_callout_slide(prs: Presentation, data: dict) -> Any:
    slide = prs.slides.add_slide(_blank_layout(prs))
    w, h = prs.slide_width, prs.slide_height
    _header(slide, data, w, h)

    categories = data.get("categories", [])
    values = data.get("values", [])
    callouts = data.get("callouts", [])

    cd = CategoryChartData()
    cd.categories = categories
    cd.add_series("Volume %", values)

    cf = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.4), Inches(1.0), Inches(4.5), Inches(5.8), cd
    )
    chart = cf.chart
    _style_chart(chart)
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(9)
    plot.data_labels.font.bold = True
    plot.data_labels.font.color.rgb = WHITE
    plot.data_labels.number_format = '0"%"'

    for series in chart.series:
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = VIOLET

    try:
        chart.category_axis.tick_labels.font.size = Pt(9)
        chart.value_axis.visible = False
        chart.value_axis.has_major_gridlines = False
    except Exception:
        pass

    card_x = Inches(5.3)
    card_w = Inches(7.7)
    n_cards = min(len(callouts), len(categories))
    if n_cards > 0:
        card_h = min(5.6 / n_cards - 0.08, 1.2)
        for ci in range(n_cards):
            cy = Inches(1.0 + ci * (card_h + 0.08))
            shape = slide.shapes.add_shape(5, card_x, cy, card_w, Inches(card_h))
            shape.fill.solid()
            shape.fill.fore_color.rgb = CARD_BG
            shape.line.color.rgb = LT_PURPLE
            shape.line.width = Pt(0.5)
            tf = shape.text_frame
            tf.word_wrap = True
            tf.margin_left = Inches(0.12)
            tf.margin_right = Inches(0.12)
            tf.margin_top = Inches(0.06)
            r = tf.paragraphs[0].add_run()
            r.text = categories[ci]
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = BRAND_VIOLET
            if ci < len(callouts) and callouts[ci]:
                p2 = tf.add_paragraph()
                r2 = p2.add_run()
                r2.text = str(callouts[ci])[:120]
                r2.font.size = Pt(8)
                r2.font.color.rgb = BLACK

    return slide


# ═══════════════════════════════════════════════════════════════════════
# 6. COLUMN CHART
# ═══════════════════════════════════════════════════════════════════════

def add_column_chart_slide(prs: Presentation, data: dict) -> Any:
    slide = prs.slides.add_slide(_blank_layout(prs))
    w, h = prs.slide_width, prs.slide_height
    _header(slide, data, w, h)

    cd = CategoryChartData()
    cd.categories = data.get("categories", [])
    cd.add_series("Volume", data.get("values", []))

    cf = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.3), Inches(1.0), Inches(12.7), Inches(4.5), cd
    )
    chart = cf.chart
    _style_chart(chart)
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(9)
    plot.data_labels.font.bold = True
    plot.data_labels.number_format = "#,##0"

    for series in chart.series:
        for pi, pt in enumerate(series.points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = SERIES_PALETTE[pi % len(SERIES_PALETTE)]

    try:
        chart.category_axis.tick_labels.font.size = Pt(9)
        chart.value_axis.visible = False
    except Exception:
        pass

    if data.get("bullets"):
        _add_bullets(slide, data["bullets"], Inches(0.4), Inches(5.7), w - Inches(0.8), Inches(1.3))

    return slide


# ═══════════════════════════════════════════════════════════════════════
# 7. PIE CHART
# ═══════════════════════════════════════════════════════════════════════

def add_pie_chart_slide(prs: Presentation, data: dict) -> Any:
    slide = prs.slides.add_slide(_blank_layout(prs))
    w, h = prs.slide_width, prs.slide_height
    _header(slide, data, w, h)

    cd = CategoryChartData()
    cd.categories = data.get("categories", [])
    cd.add_series("Share", data.get("values", []))

    cf = slide.shapes.add_chart(
        XL_CHART_TYPE.PIE, Inches(0.4), Inches(1.0), Inches(6.0), Inches(5.5), cd
    )
    chart = cf.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.font.size = Pt(9)
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(10)
    plot.data_labels.font.bold = True
    plot.data_labels.font.color.rgb = WHITE
    plot.data_labels.number_format = '0"%"'

    for pi, pt in enumerate(chart.series[0].points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = SERIES_PALETTE[pi % len(SERIES_PALETTE)]

    if data.get("bullets"):
        _add_bullets(slide, data["bullets"], Inches(7.0), Inches(1.2), Inches(5.8), Inches(5.0))

    return slide


# ═══════════════════════════════════════════════════════════════════════
# 8. STACKED BAR 100%
# ═══════════════════════════════════════════════════════════════════════

def add_stacked_bar_slide(prs: Presentation, data: dict) -> Any:
    slide = prs.slides.add_slide(_blank_layout(prs))
    w, h = prs.slide_width, prs.slide_height
    _header(slide, data, w, h)

    cd = CategoryChartData()
    cd.categories = data.get("categories", [])
    cd.add_series("Positive", data.get("positive", []))
    cd.add_series("Negative", data.get("negative", []))
    cd.add_series("Neutral", data.get("neutral", []))

    cf = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_STACKED_100,
        Inches(0.3), Inches(1.0), Inches(12.7), Inches(5.0), cd
    )
    chart = cf.chart
    _style_chart(chart, legend=True)

    for si, color in enumerate([VIOLET, RED, BLUE]):
        chart.series[si].format.fill.solid()
        chart.series[si].format.fill.fore_color.rgb = color

    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(8)
    plot.data_labels.font.color.rgb = WHITE
    plot.data_labels.number_format = '0"%"'

    try:
        chart.category_axis.tick_labels.font.size = Pt(9)
        chart.value_axis.visible = False
    except Exception:
        pass

    if data.get("bullets"):
        _add_bullets(slide, data["bullets"], Inches(0.4), Inches(6.3), w - Inches(0.8), Inches(0.8))

    return slide


# ═══════════════════════════════════════════════════════════════════════
# 9. COMPARISON TABLE
# ═══════════════════════════════════════════════════════════════════════

def add_comparison_table_slide(prs: Presentation, data: dict) -> Any:
    slide = prs.slides.add_slide(_blank_layout(prs))
    w, h = prs.slide_width, prs.slide_height
    _header(slide, data, w, h)

    headers = data.get("headers", [])
    rows = data.get("rows", [])
    n_cols = len(headers) if headers else (len(rows[0]) if rows else 3)
    n_rows = len(rows) + 1

    tbl_w = Inches(12.5)
    tbl_h = Inches(min(n_rows * 0.7, 5.8))
    table_shape = slide.shapes.add_table(n_rows, n_cols, Inches(0.4), Inches(0.9), tbl_w, tbl_h)
    table = table_shape.table

    for ci, header in enumerate(headers):
        cell = table.cell(0, ci)
        cell.text = str(header)
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = WHITE
        cell.fill.solid()
        cell.fill.fore_color.rgb = BRAND_VIOLET

    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri + 1, ci)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(9)
                p.font.color.rgb = BLACK
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if ri % 2 == 0 else WHITE
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    return slide


# ═══════════════════════════════════════════════════════════════════════
# 10. THEME MATRIX — 2-column cards with gauges
# ═══════════════════════════════════════════════════════════════════════

def add_theme_matrix_slide(prs: Presentation, data: dict) -> Any:
    slide = prs.slides.add_slide(_blank_layout(prs))
    w, h = prs.slide_width, prs.slide_height
    _header(slide, data, w, h)

    lh = data.get("left_header", "KEY FINDINGS")
    rh = data.get("right_header", "IMPLICATIONS")

    for hdr_text, hdr_x in [(lh, Inches(0.4)), (rh, Inches(7.0))]:
        box = slide.shapes.add_textbox(hdr_x, Inches(0.85), Inches(5.5), Inches(0.3))
        r = box.text_frame.paragraphs[0].add_run()
        r.text = hdr_text
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = BRAND_VIOLET

    items = data.get("items", [])
    n = len(items)
    if n == 0:
        return slide

    item_h = min(5.2 / n - 0.08, 1.2)

    for idx, item in enumerate(items):
        y = Inches(1.2 + idx * (item_h + 0.08))

        if item.get("pct") is not None:
            gd = CategoryChartData()
            gd.categories = ["Fill", "Empty"]
            gd.add_series("G", (item["pct"], 100 - item["pct"]))
            gf = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(0.3), y, Inches(0.7), Inches(item_h * 0.7), gd)
            gc = gf.chart
            _style_chart(gc)
            gc.plots[0].has_data_labels = False
            gc.series[0].points[0].format.fill.solid()
            gc.series[0].points[0].format.fill.fore_color.rgb = VIOLET
            gc.series[0].points[1].format.fill.solid()
            gc.series[0].points[1].format.fill.fore_color.rgb = LT_PURPLE
            left_x = Inches(1.1)
        else:
            left_x = Inches(0.4)

        for col_x, title_key, desc_key in [
            (left_x, "left_title", "left_desc"),
            (Inches(7.0), "right_title", "right_desc"),
        ]:
            col_w = Inches(5.8) if col_x < Inches(5) else Inches(5.5)
            tb = slide.shapes.add_textbox(col_x, y, col_w, Inches(0.22))
            r = tb.text_frame.paragraphs[0].add_run()
            r.text = item.get(title_key, "")
            r.font.size = Pt(9)
            r.font.bold = True
            r.font.color.rgb = BLACK

            db = slide.shapes.add_textbox(col_x, y + Inches(0.22), col_w, Inches(item_h - 0.25))
            db.text_frame.word_wrap = True
            r = db.text_frame.paragraphs[0].add_run()
            r.text = item.get(desc_key, "")[:100]
            r.font.size = Pt(8)
            r.font.color.rgb = GREY

    return slide


# ═══════════════════════════════════════════════════════════════════════
# 11. TWO-BY-TWO FRAMEWORK
# ═══════════════════════════════════════════════════════════════════════

def add_two_by_two_slide(prs: Presentation, data: dict) -> Any:
    """2x2 strategic framework (e.g., opportunity vs risk, high/low axes).

    data: x_label, y_label,
          quadrants: [{title, items: [str, ...]}, ...] (TL, TR, BL, BR)
    """
    slide = prs.slides.add_slide(_blank_layout(prs))
    w, h = prs.slide_width, prs.slide_height
    _header(slide, data, w, h)

    x_label = data.get("x_label", "")
    y_label = data.get("y_label", "")
    quadrants = data.get("quadrants", [])

    grid_left = Inches(1.5)
    grid_top = Inches(1.2)
    grid_w = Inches(10.5)
    grid_h = Inches(5.2)
    half_w = grid_w / 2
    half_h = grid_h / 2

    colors = [
        RGBColor(0xE8, 0xDE, 0xF8),
        RGBColor(0xD1, 0xC4, 0xE9),
        RGBColor(0xF3, 0xE5, 0xF5),
        RGBColor(0xED, 0xE7, 0xF6),
    ]

    positions = [
        (grid_left, grid_top),
        (grid_left + half_w, grid_top),
        (grid_left, grid_top + half_h),
        (grid_left + half_w, grid_top + half_h),
    ]

    for qi, (qx, qy) in enumerate(positions):
        shape = slide.shapes.add_shape(1, qx, qy, half_w, half_h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = colors[qi % 4]
        shape.line.color.rgb = BRAND_VIOLET
        shape.line.width = Pt(0.5)

        if qi < len(quadrants):
            q = quadrants[qi]
            tf = shape.text_frame
            tf.word_wrap = True
            tf.margin_left = Inches(0.15)
            tf.margin_top = Inches(0.15)

            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = q.get("title", "")
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = BRAND_VIOLET

            for item_text in q.get("items", [])[:4]:
                p2 = tf.add_paragraph()
                r2 = p2.add_run()
                r2.text = f"• {item_text}"
                r2.font.size = Pt(8)
                r2.font.color.rgb = BLACK
                p2.space_before = Pt(3)

    yb = slide.shapes.add_textbox(Inches(0.2), grid_top + grid_h / 2 - Inches(0.15), Inches(1.2), Inches(0.3))
    yb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = yb.text_frame.paragraphs[0].add_run()
    r.text = y_label
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = BRAND_VIOLET

    xb = slide.shapes.add_textbox(grid_left + grid_w / 2 - Inches(1), grid_top + grid_h + Inches(0.05), Inches(2), Inches(0.25))
    xb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = xb.text_frame.paragraphs[0].add_run()
    r.text = x_label
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = BRAND_VIOLET

    return slide


# ═══════════════════════════════════════════════════════════════════════
# DISPATCHER
# ═══════════════════════════════════════════════════════════════════════

VISUAL_BUILDERS = {
    "kpi_tiles": add_kpi_tiles_slide,
    "doughnut_sentiment": add_doughnut_sentiment_slide,
    "doughnut_gauge": add_doughnut_gauge_slide,
    "line_trend": add_line_trend_slide,
    "bar_callout": add_bar_callout_slide,
    "column_chart": add_column_chart_slide,
    "pie_chart": add_pie_chart_slide,
    "stacked_bar": add_stacked_bar_slide,
    "comparison_table": add_comparison_table_slide,
    "theme_matrix": add_theme_matrix_slide,
    "two_by_two": add_two_by_two_slide,
}


def build_visual_slide(prs: Presentation, slide_data: dict) -> Any:
    vtype = slide_data.get("visual_type", "")
    builder = VISUAL_BUILDERS.get(vtype)
    if builder:
        return builder(prs, slide_data)
    return None
