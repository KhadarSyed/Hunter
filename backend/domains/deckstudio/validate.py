"""Deck validator: after the deck is built, every requirement of a client deck is checked and each gap named --
the client's logo and product on the cover, the brand vs competitor chapter, an answer, a chart, a cited summary
and verbatims for every question, a cited summary on every chart slide, citations that name a brand literally (not
an idiom like "a band-aid solution") and are listed (with their icon) in Sources, competitors' verbatims on their own
slide with the post or article screenshot, icons beside chart labels, the social vs editorial split, slides that
build on click in PowerPoint, and the brief's own asks.

Rows are {check, status: pass | warn | fail, detail}; a fail is something the client would notice is missing.
"""
from __future__ import annotations

import re
from pathlib import Path

from ..deliverable.engine_types import RQ
from ..deliverable.insights import _off_segment
from ..deliverable.visuals import is_badge, value_icon
from .spec import DeckSpec, SlideSpec


def _row(check: str, ok: bool, detail: str = "", soft: bool = False, note: str = "") -> dict:
    """`note`: what a passing check still wants the analyst to know (e.g. a category shown as a count only)."""
    return {"check": check, "status": "pass" if ok else ("warn" if soft else "fail"), "detail": note if ok else detail}


def _cited(slides: list[SlideSpec]) -> set[str]:
    return {str(n) for s in slides for c in s.cards for n in c.get("citations") or []}


def _question(rq: RQ, slides: list[SlideSpec]) -> dict:
    mine = [s for s in slides if s.id.startswith(rq.id.lower() + "-")]
    divider = next((s for s in mine if s.type == "divider"), None)
    gaps = []
    if not divider or not (divider.so_what or "").strip():
        gaps.append("no answer")
    if not any(s.charts or s.tables for s in mine if s.type != "verbatim_wall"):
        gaps.append("no chart")
    if not _cited([s for s in mine if s.type != "verbatim_wall"]):
        gaps.append("no cited summary")
    if not any(s.type == "verbatim_wall" and s.cards for s in mine):
        gaps.append("no verbatims")
    return _row(rq.question, not gaps, ", ".join(gaps))


def _chart_slides(slides: list[SlideSpec], rqs: list[RQ]) -> list[SlideSpec]:
    """The evidence slides: a chart or table in the brand chapter or a question's chapter."""
    prefixes = ("brand-",) + tuple(rq.id.lower() + "-" for rq in rqs)
    return [s for s in slides if s.type not in ("verbatim_wall", "divider") and (s.charts or s.tables)
            and s.id.startswith(prefixes)]


# charts of things that have a logo or site icon; people (experts, celebrities) and Yes/No shares have none
_ENTITY_CHART = re.compile(r"brand[-_]sov|brand[-_]sentiment|outlet_ranking|entities-brands|entities-retailers")
# charts of brands and retailers, whose mark should be the verified logo (an outlet without a site icon shows its
# initial by design: a court-records feed or a broadcast clip has no logo to verify)
_BRAND_CHART = re.compile(r"brand[-_]sov|brand[-_]sentiment|entities-brands|entities-retailers")
# categories that answer nothing on their own; a chart of only these is a placeholder, and they get no insight card
_NO_ANSWER = {"unknown", "none", "other", "not stated", "n/a"}
SEGMENT_CARDS = 5                  # matches the planner: a breakdown chart explains up to 5 of its categories


def _segments(slide: SlideSpec) -> list[str]:
    """The categories of a breakdown chart that each need their own explanation card."""
    chart = slide.charts[0] if slide.charts else {}
    if chart.get("kind") not in ("bar", "treemap", "doughnut", "pie"):
        return []
    cats = [str(c) for c in chart.get("categories") or [] if str(c).strip().lower() not in _NO_ANSWER]
    return cats[:SEGMENT_CARDS] if len(cats) >= 2 else []
_NAMED_THINGS = re.compile(r"brand|competitor|retailer|product|outlet|expert|celebrit|people|person|entities")


def off_topic_cards(charts: list[SlideSpec]) -> list[tuple[str, str]]:
    """(slide id, card headline) for each category card whose text never names its category."""
    # brand, outlet and people charts are left out: their cards name brands and people by design, and a brand's
    # name in a theme card is not a sign the card is about the wrong theme
    charts = [s for s in charts if not _NAMED_THINGS.search(s.id)]
    out = []
    by_question: dict[str, set[str]] = {}
    for s in charts:
        by_question.setdefault(s.id.split("-", 1)[0], set()).update(g.strip().lower() for g in _segments(s))
    for s in charts:
        segs = {g.strip().lower(): {0} for g in _segments(s)}
        related = by_question.get(s.id.split("-", 1)[0], set())
        for c in s.cards:
            key = str(c.get("headline", "")).strip().rstrip(":").lower()
            if key in segs and not c.get("count_only") and _off_segment(str(c.get("headline")), str(c.get("text", "")), segs,
                                                                     f"{s.question} {s.title} {s.kicker}", related):
                out.append((s.id, str(c.get("headline"))))
    return out


_STATIC = {"objectives", "methodology", "checklist", "citations", "scorecard"}   # reading slides: nothing to build


TERRITORY_SHARE = 0.2      # under this share of articles naming the brand, the study is about a territory


def _evidence_rows(slides: list[SlideSpec], rqs: list[RQ], has_competitors: bool,
                   citation_brands: dict[str, list[str]] | None, mixed_media: bool = False,
                   territory: bool = False) -> list[dict]:
    charts = _chart_slides(slides, rqs)
    bare = [s.id for s in charts if not _cited([s]) and "-cont-" not in s.id]   # a continued table's summary is on its first part
    rows = [_row("Every chart slide has a cited summary", not bare, "no cited summary on " + ", ".join(bare[:6]))]
    # icons belong beside named things (brands, outlets, retailers, experts, celebrities), not Yes/No or months
    unlabelled = [s.id for s in charts
                  if (cats := [c for ch in s.charts for c in ch.get("categories") or []]) and not set(cats) & set(s.logos)
                  and any(value_icon(str(c)) for c in cats) and not _ENTITY_CHART.search(s.id)]
    rows.append(_row("Chart labels show their icon or logo", not unlabelled,
                     f"{len(unlabelled)} chart(s) without icons: " + ", ".join(unlabelled[:4]), soft=True))
    # a brand, retailer or outlet chart marks every label: one logo among eight bars reads as a broken slide
    no_logo = [f"{s.id} ({', '.join(missing[:4])})" for s in charts if _ENTITY_CHART.search(s.id)
               for missing in [[str(c) for ch in s.charts for c in ch.get("categories") or [] if str(c) not in s.logos]]
               if missing]
    rows.append(_row("Every brand, retailer and outlet in a chart shows its logo", not no_logo,
                     "missing: " + "; ".join(no_logo[:4])))
    # a letter badge keeps the chart readable but is not the brand's logo: name every label still waiting for one
    # (soft: some brands have no logo that the vision check could confirm, and a wrong logo is worse than a badge)
    badged = sorted({str(c) for s in charts if _BRAND_CHART.search(s.id) for ch in s.charts
                     for c in ch.get("categories") or [] if str(c) in s.logos and is_badge(s.logos[str(c)])})
    rows.append(_row("Brand logos are verified, not letter-badge stand-ins", not badged,
                     f"{len(badged)} without a verified logo: " + ", ".join(badged[:10]), soft=True))
    # each category of a breakdown chart is explained on its own card (what drives it), not one summary for all
    thin = [f"{s.id} ({len(s.cards)} of {len(segs)})" for s in charts if (segs := _segments(s)) and len(s.cards) < len(segs)]
    counted = [f"{c['headline']} ({s.id})" for s in charts for c in s.cards if c.get("count_only")]
    rows.append(_row("Every chart category has its own insight", not thin,
                     "explained only partly: " + ", ".join(thin[:6]),
                     note=f"{len(counted)} category(ies) shown as a count only - no sampled article explains them: "
                          + ", ".join(counted[:6]) if counted else ""))
    # a category's card is about that category: an 'expansion' card that writes about a scandal explains the wrong bar
    off = off_topic_cards(charts)
    rows.append(_row("Each category card explains its own category", not off,
                     "card(s) about something else: " + ", ".join(f"{h} ({sid})" for sid, h in off[:6])))
    empty = [s.id for s in charts for ch in s.charts[:1]
             if ch.get("categories") and all(str(c).strip().lower() in _NO_ANSWER for c in ch["categories"])]
    rows.append(_row("No chart shows only 'Not stated' or 'Unknown'", not empty, "placeholder chart: " + ", ".join(empty)))
    if citation_brands is not None:
        stray = sorted((n for n in _cited(slides) if not citation_brands.get(n)), key=lambda n: int(n) if n.isdigit() else 0)
        # a territory study (sports-injury talk for a first-aid brand) cites the conversation itself: there a source
        # that never names the brand is the expected evidence, so the check passes with a note
        rows.append(_row("Citations name the brand or a competitor literally", not stray or territory,
                         "cited source(s) never name a brand: " + ", ".join(stray[:12]),
                         note=f"{len(stray)} cited source(s) discuss the wider conversation without naming a brand, "
                              "as expected for a territory study" if stray and territory else ""))
    walls = [s for s in slides if s.type == "verbatim_wall"]
    if has_competitors:
        rows.append(_row("Competitor verbatims on their own slide",
                         any(s.id.endswith("competitors-verbatims") and s.cards for s in walls),
                         "no competitor verbatim slide", soft=True))
    no_shot = sum(1 for s in walls for c in s.cards if not c.get("image"))
    rows.append(_row("Verbatims show the post or article screenshot", not no_shot,
                     f"{no_shot} verbatim(s) without a screenshot", soft=True))
    if mixed_media:                       # a single-channel study (editorial only, or social only) has nothing to split
        rows.append(_row("Social vs editorial split", any("social vs editorial" in (s.title or "").lower() for s in charts),
                         "no social vs editorial chart", soft=True))
    return rows


def _builds_row(spec: DeckSpec, pptx: Path) -> dict:
    """Slides with cards or verbatims must reveal them on click (a <p:timing> in the slide XML)."""
    from pptx import Presentation
    ns = "{http://schemas.openxmlformats.org/presentationml/2006/main}timing"
    try:
        animated = [s._element.find(ns) is not None for s in Presentation(str(pptx)).slides]
    except Exception as e:                  # an unreadable file is itself a failure to report, not to raise
        return _row("Slides build on click in PowerPoint", False, f"could not read the PPTX: {type(e).__name__}")
    wanted = [k for k, s in enumerate(spec.slides) if s.cards and s.type not in _STATIC and k < len(animated)]
    still = [spec.slides[k].id for k in wanted if not animated[k]]
    if wanted and len(still) == len(wanted):
        return _row("Slides build on click in PowerPoint", False, "no slide animates its cards")
    return _row("Slides build on click in PowerPoint", not still, f"{len(still)} slide(s) without builds: "
                + ", ".join(still[:4]), soft=True)


def validate_deck(spec: DeckSpec, rqs: list[RQ], checklist: list[dict], brand_products: list[str],
                  has_competitors: bool, citation_brands: dict[str, list[str]] | None = None,
                  pptx: Path | None = None, is_category: bool = False, mixed_media: bool = False,
                  brand_share: float = 1.0) -> list[dict]:
    """`is_category`: the client is a category, not a brand -- no single logo, and coverage needn't name a brand."""
    slides = spec.slides
    rows = [_row("Client logo on the cover", bool(spec.brand_logo),
                 "a category study has no single brand logo" if is_category else "no logo for the client brand",
                 soft=is_category),
            _row("Client product on the cover", bool(spec.product_image) or not brand_products,
                 f"no photo found for {', '.join(brand_products[:2])}", soft=True)]
    if has_competitors:
        rows.append(_row("Brand share of voice vs competitors", any(s.id.startswith("brand-sov") for s in slides),
                         "no share of voice slide"))
        rows.append(_row("Brand sentiment vs competitors", any(s.id.startswith("brand-sentiment") for s in slides),
                         "no sentiment slide"))
    rows += [_question(rq, slides) for rq in rqs]
    rows += _evidence_rows(slides, rqs, has_competitors, None if is_category else citation_brands, mixed_media,
                           territory=brand_share < TERRITORY_SHARE)
    summary = next((s for s in slides if s.type == "executive_summary"), None)
    uncited = [c.get("headline") or "" for c in (summary.cards if summary else []) if not c.get("citations")]
    rows.append(_row("Executive summary cites its sources", bool(summary) and not uncited,
                     f"{len(uncited)} answer(s) without a source" if summary else "no executive summary", soft=True))
    listed = {str(r[0]) for s in slides if s.type == "citations" for t in s.tables for r in t.get("rows", []) if r}
    missing = sorted(_cited(slides) - listed, key=lambda n: int(n) if n.isdigit() else 0)
    rows.append(_row("Every citation is listed in Sources", not missing, "not in Sources: " + ", ".join(missing)))
    no_icon = sorted(n for n in _cited(slides) if not (spec.citation_meta.get(n) or {}).get("icon"))
    rows.append(_row("Citations show their source icon", not no_icon, f"{len(no_icon)} source(s) without an icon",
                     soft=True))
    if pptx:
        rows.append(_builds_row(spec, pptx))
    rows += [_row(f"Brief: {r['ask']}", r["status"] == "covered", r.get("note") or r["status"],
                  soft=r["status"] == "partial")
             for r in checklist if r.get("kind") != "question"]
    return rows
